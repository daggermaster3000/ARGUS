"""End-to-end GUI check: build the viewer, load samples, measure, export.

Runs against a real Qt application but never needs a visible screen::

    python tests/smoke_gui.py

Set ``QT_QPA_PLATFORM=offscreen`` to run it on a machine with no display.
"""

from __future__ import annotations

import shutil
import sys
import tempfile
import time
from pathlib import Path

import numpy as np
from qtpy.QtCore import QCoreApplication, Qt

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

SAMPLES = Path(__file__).resolve().parent / "sample_data"

_failures: list[str] = []


#: Signatures of "this machine has no usable OpenGL context right now": the
#: offscreen Qt platform, and a Remote Desktop session that has been disconnected
#: (which tears down the session's GPU context under a running process).
_NO_GL_MARKERS = (
    "no OpenGL context",
    "Failed to create context",
    "'NoneType' object has no attribute 'setsize'",
)


def _is_missing_gl(exc: BaseException) -> bool:
    return any(marker in str(exc) for marker in _NO_GL_MARKERS)


def load_and_release(path):
    """Read a file and give its handle straight back.

    The readers hold files open for the life of the process so their lazy arrays
    stay readable; in a temporary folder that is what stops Windows deleting it.
    """
    from microscopy_viewer.loaders import load_path, release

    specs = load_path(path)
    for spec in specs:
        spec.data = np.asarray(spec.data[0] if spec.multiscale else spec.data)
        spec.multiscale = False
    release(path)
    return specs


def loader_release(path):
    from microscopy_viewer.loaders import release

    return release(path)


def check(condition: bool, message: str) -> None:
    if condition:
        print(f"  ok   {message}", flush=True)
    else:
        print(f"  FAIL {message}", flush=True)
        _failures.append(message)



class _StubSegmentation:
    """A segmentation backend that labels a corner, so no weights are downloaded."""

    name = "stub"
    install_hint = "-"

    def available(self) -> bool:
        return True

    def model_choices(self):
        from microscopy_viewer import segmentation as sg

        return (sg.ModelChoice(label="stub", value="stub"),)

    def models(self):
        return ("stub",)

    def available_models(self):
        return ("stub",)

    def unsupported_models(self):
        return ()

    def segment(self, image, settings, diameter_px, anisotropy, device, progress=None):
        import numpy as np

        masks = np.zeros(np.asarray(image).shape, dtype=np.int32)
        masks[..., :4, :4] = 1
        return masks, {}


def main() -> int:
    if not SAMPLES.exists():
        print("sample data missing — run tests/make_sample_data.py first")
        return 2

    from microscopy_viewer import measurements as mm
    from microscopy_viewer.app import MicroscopyViewer
    from microscopy_viewer.contrast import auto_contrast, reset_contrast
    from microscopy_viewer.exports import export_snapshot
    from microscopy_viewer.utils import MICRON

    print("building the viewer", flush=True)
    app = MicroscopyViewer(show=False)
    check(app.viewer is not None, "napari viewer created")
    check(app.toolbar is not None, "toolbar built")
    check(app.metadata_widget is not None, "metadata dock built")
    check(app.measurements_widget is not None, "measurements dock built")
    check(app._drop_filter is not None, "drag-and-drop filter installed")
    print(flush=True)

    print("opening samples", flush=True)
    paths = [
        SAMPLES / "sample_4d_2ch.ims",
        SAMPLES / "sample_tzcyx.ome.tif",
        SAMPLES / "sample_plain.tif",
    ]
    added = app.open_paths(paths)
    check(added == 5, f"five layers added from three files (got {added})")
    check(len(app.viewer.layers) == 5, f"viewer holds five layers ({len(app.viewer.layers)})")
    check(app.viewer.dims.ndim == 4, f"viewer is 4D for the widest dataset ({app.viewer.dims.ndim})")
    # Where the scale bar reads its unit from moved in napari 0.8: the overlay
    # lost its own ``unit`` field and now takes it from the layer. Check whichever
    # one this napari actually has, because checking the wrong one passes on the
    # version that is broken and fails on the version that works.
    overlay_fields = getattr(type(app.viewer.scale_bar), "model_fields", None) or getattr(
        type(app.viewer.scale_bar), "__fields__", {}
    )
    if "unit" in overlay_fields:
        check(app.viewer.scale_bar.unit == MICRON, f"scale bar unit = {app.viewer.scale_bar.unit!r}")
    else:
        # napari normalises the unit through pint, so "um" comes back as
        # "micrometer" rather than as the string the reader handed it.
        micron_names = {MICRON, "um", "micron", "micrometer", "micrometre"}
        units = [tuple(str(u) for u in layer.units) for layer in app.viewer.layers]
        check(
            all(unit in micron_names for entry in units for unit in entry[-2:]),
            f"layers carry micrometre units for the scale bar to read ({units[0]})",
        )
    names = [layer.name for layer in app.viewer.layers]
    check(any("GFP" in name for name in names), f"channel names used in layers: {names[:2]}")
    print(flush=True)

    print("metadata panel follows the active layer", flush=True)
    ims_layer = app.viewer.layers[0]
    app.viewer.layers.selection.active = ims_layer
    app.metadata_widget.refresh()
    tree = app.metadata_widget._tree
    check(tree.topLevelItemCount() >= 2, f"tree has sections ({tree.topLevelItemCount()})")
    labels = [tree.topLevelItem(i).text(0) for i in range(tree.topLevelItemCount())]
    check(any("Channel" in label for label in labels), f"a channel section is present: {labels}")

    rows = {}
    dataset = tree.topLevelItem(0)
    for index in range(dataset.childCount()):
        rows[dataset.child(index).text(0)] = dataset.child(index).text(1)
    check("Pixel size" in rows and "0.13" in rows["Pixel size"], f"pixel size row = {rows.get('Pixel size')!r}")
    check("Z-step size" in rows and "0.5" in rows["Z-step size"], f"z-step row = {rows.get('Z-step size')!r}")
    check("Objective" in rows, f"objective row = {rows.get('Objective')!r}")
    check("Numerical aperture" in rows, f"NA row = {rows.get('Numerical aperture')!r}")

    tiff_layer = app.viewer.layers[2]
    app.viewer.layers.selection.active = tiff_layer
    heading_before = app.metadata_widget._heading.text()
    app.viewer.layers.selection.active = ims_layer
    heading_after = app.metadata_widget._heading.text()
    check(heading_before != heading_after, "heading changes when the active layer changes")
    print(flush=True)

    print("contrast controls", flush=True)
    limits_before = tuple(ims_layer.contrast_limits)
    ims_layer.contrast_limits = (0, 1)
    changed = auto_contrast(app.viewer)
    check(changed >= 1, f"auto contrast touched {changed} layer(s)")
    check(tuple(ims_layer.contrast_limits) != (0, 1), f"limits updated to {tuple(ims_layer.contrast_limits)}")
    reset = reset_contrast(app.viewer)
    check(reset == 5, f"reset contrast touched every image layer ({reset})")
    check(limits_before is not None, "original limits captured")
    print(flush=True)

    print("ROI drawing and measurement", flush=True)
    app.viewer.layers.selection.active = ims_layer
    app.measurements_widget.new_roi_layer()
    shapes = mm.shapes_layers(app.viewer)
    check(len(shapes) == 1, f"one ROI layer created ({len(shapes)})")
    roi = shapes[0]
    check(roi.ndim == ims_layer.ndim, f"ROI layer matches image ndim ({roi.ndim})")
    check(tuple(roi.scale) == tuple(ims_layer.scale), f"ROI layer inherits scale {tuple(roi.scale)}")

    # A 20 x 40 px rectangle on t=0, z=2: 0.13 µm pixels -> 2.6 x 5.2 µm = 13.52 µm².
    rectangle = np.array(
        [[0, 2, 0, 0], [0, 2, 20, 0], [0, 2, 20, 40], [0, 2, 0, 40]], dtype=float
    )
    roi.add_rectangles(rectangle)
    mm.ensure_name_feature(roi)
    app.measurements_widget.recompute()
    results = {m.measurement_type: m for m in app.measurements_widget._rows}
    check("Area" in results, f"an area was measured ({list(results)})")
    if "Area" in results:
        expected = 20 * 0.13 * 40 * 0.13
        check(
            abs(results["Area"].value - expected) < 1e-6,
            f"area = {results['Area'].value:.4f} µm² (expected {expected:.4f})",
        )
        check(results["Area"].unit == "µm²", f"unit = {results['Area'].unit}")
        check(results["Area"].image_name == ims_layer.name, f"image = {results['Area'].image_name}")
        check(results["Area"].channel != "", f"channel = {results['Area'].channel!r}")
        check("Z=2" in results["Area"].slice_position, f"slice = {results['Area'].slice_position!r}")
    check(app.measurements_widget._table.rowCount() == len(app.measurements_widget._rows), "table matches results")

    mm.rename_roi(roi, 0, "Nucleus 1")
    app.measurements_widget.recompute()
    check(
        all(m.roi_name == "Nucleus 1" for m in app.measurements_widget._rows),
        f"rename propagated: {[m.roi_name for m in app.measurements_widget._rows]}",
    )
    print(flush=True)

    print("3D / MIP shows real resolution", flush=True)
    from microscopy_viewer.rendering import _LEVEL, _PYRAMID

    # ims_layer is a 2-level Imaris pyramid: 128x160 at level 0, 64x80 at level 1.
    check(ims_layer.multiscale, "the Imaris layer is multiscale")
    check(len(ims_layer.data) == 2, f"two levels in 2D ({len(ims_layer.data)})")
    world_2d = ims_layer.extent.world.copy()
    scale_2d = tuple(ims_layer.scale)

    app.toolbar.toggle_ndisplay()
    check(app.viewer.dims.ndisplay == 3, f"viewer switched to 3D ({app.viewer.dims.ndisplay})")
    check(ims_layer.rendering == "mip", f"rendering set to MIP ({ims_layer.rendering})")
    # Without the fix napari would render len(data)-1, i.e. the 64x80 level.
    check(len(ims_layer.data) == 1, f"pyramid collapsed to one level ({len(ims_layer.data)})")
    check(
        tuple(ims_layer.data[0].shape[-2:]) == (128, 160),
        f"3D shows full resolution {tuple(ims_layer.data[0].shape[-2:])}, not the 64x80 level",
    )
    check(ims_layer.metadata[_LEVEL] == 0, f"level 0 chosen ({ims_layer.metadata[_LEVEL]})")
    check(tuple(ims_layer.scale) == scale_2d, "scale unchanged when level 0 is used")

    app.toolbar.toggle_ndisplay()
    check(app.viewer.dims.ndisplay == 2, "viewer switched back to 2D")
    check(len(ims_layer.data) == 2, f"full pyramid restored ({len(ims_layer.data)})")
    check(ims_layer.metadata[_LEVEL] is None, "layer marked as back on the full pyramid")
    check(tuple(ims_layer.scale) == scale_2d, "scale restored")
    check(np.allclose(ims_layer.extent.world, world_2d), "world extent unchanged across the round trip")

    # Force a budget so small that the coarse level must be used, and confirm the
    # scale compensates so world coordinates still line up.
    from microscopy_viewer.rendering import MultiscaleDepthManager

    tight = MultiscaleDepthManager(app.viewer, budget=1000)
    app.viewer.dims.ndisplay = 3
    tight.apply()
    check(ims_layer.metadata[_LEVEL] == 1, f"tight budget steps down to level 1 ({ims_layer.metadata[_LEVEL]})")
    check(tuple(ims_layer.data[0].shape[-2:]) == (64, 80), f"coarse level in use {tuple(ims_layer.data[0].shape[-2:])}")
    check(
        np.allclose(np.array(ims_layer.scale)[-2:], np.array(scale_2d)[-2:] * 2),
        f"voxel size doubled to match {tuple(ims_layer.scale)[-2:]}",
    )
    app.viewer.dims.ndisplay = 2
    tight.apply()
    check(len(ims_layer.data) == 2 and tuple(ims_layer.scale) == scale_2d, "restored after the tight-budget run")
    check(_PYRAMID in ims_layer.metadata, "the original pyramid is still held on the layer")
    print(flush=True)

    print("exports", flush=True)
    with tempfile.TemporaryDirectory() as directory:
        out = Path(directory)
        # Canvas screenshots need a live OpenGL context. That is absent under
        # QT_QPA_PLATFORM=offscreen and on a Remote Desktop session that has been
        # disconnected, so treat it as a skip rather than failing the whole run.
        try:
            png = export_snapshot(app.viewer, out / "shot.png", scale=1.0)
            gl_available = True
        except Exception as exc:
            if not _is_missing_gl(exc):
                raise
            gl_available = False
            print(f"  skip snapshot checks — no OpenGL context ({exc})", flush=True)

        if gl_available:
            check(png.exists() and png.stat().st_size > 1000, f"PNG snapshot written ({png.stat().st_size} bytes)")
            tif = export_snapshot(app.viewer, out / "shot.tif", scale=1.0)
            check(tif.exists() and tif.stat().st_size > 1000, f"TIFF snapshot written ({tif.stat().st_size} bytes)")

            import tifffile

            with tifffile.TiffFile(tif) as handle:
                check(handle.pages[0].shape[0] > 0, f"TIFF snapshot is a real image {handle.pages[0].shape}")

        from microscopy_viewer.exports import export_measurements

        book = export_measurements(
            app.measurements_widget._rows,
            out / "measurements.xlsx",
            [layer.metadata.get("mv_metadata") for layer in app.viewer.layers],
        )
        import openpyxl

        loaded = openpyxl.load_workbook(book)
        sheet = loaded["Measurements"]
        check(sheet.max_row == len(app.measurements_widget._rows) + 1, f"workbook has a header + data rows ({sheet.max_row})")
        acquisition = loaded["Acquisition"]
        check(acquisition.max_row > 5, f"acquisition sheet populated ({acquisition.max_row} rows)")
    print(flush=True)

    print("time series playback", flush=True)
    from microscopy_viewer import movie, timeseries
    from microscopy_viewer import volume_cache

    player = app.timeseries_widget
    check(player is not None, "the time-series panel was built")
    axis = timeseries.viewer_time_axis(app.viewer)
    check(axis == 0, f"the time axis is the first slider ({axis})")
    # The sample .ims has three timepoints; the OME-TIFF beside it has more, and
    # napari's slider spans the widest dataset, so the count is read rather than
    # assumed.
    count = player.count
    check(count >= 3, f"the time slider spans {count} timepoints")
    player.use_full_range()
    check(player.range() == (0, count - 1), f"Full range covers the whole series {player.range()}")

    timeseries.set_index(app.viewer, 0, axis)
    player.step_forward()
    check(timeseries.current_index(app.viewer, axis) == 1, "the step button moves one timepoint")
    player.go_last()
    check(timeseries.current_index(app.viewer, axis) == count - 1, "and the end button jumps to the last one")
    player.step_forward()
    check(timeseries.current_index(app.viewer, axis) == 0, "stepping past the end wraps round")

    player._fps.setValue(30.0)
    player.play()
    check(player.playing, "playback started")
    from qtpy.QtWidgets import QApplication

    started = timeseries.current_index(app.viewer, axis)
    deadline = time.monotonic() + 3.0
    while timeseries.current_index(app.viewer, axis) == started and time.monotonic() < deadline:
        QApplication.instance().processEvents()
        time.sleep(0.01)
    check(
        timeseries.current_index(app.viewer, axis) != started,
        "and the clock advanced the slider without anyone touching it",
    )
    player.pause()
    check(not player.playing, "pause stops the timer")

    # The sample data is far below the cache's size floor, so lower it to check
    # the copy actually happens rather than being skipped as not worth it.
    ims_timeline = next(layer for layer in app.viewer.layers if timeseries.has_timeline(layer))
    floor = volume_cache.MIN_CACHE_BYTES
    volume_cache.MIN_CACHE_BYTES = 0
    try:
        before_shapes = [tuple(level.shape) for level in ims_timeline.data]
        before_scale = tuple(ims_timeline.scale)
        local = timeseries.TimelineManager(app.viewer, threaded=False)
        local.cache_layer(ims_timeline, force=True)
        check(local.is_cached(ims_timeline), "the time series was copied to the local cache")
        check(
            [tuple(level.shape) for level in ims_timeline.data] == before_shapes,
            "the cached levels have the shapes they replaced, so nothing about the view changed",
        )
        check(tuple(ims_timeline.scale) == before_scale, "and the voxel size is untouched")
        check(
            "local" in local.describe(ims_timeline),
            f"the panel reports where frames come from ({local.describe(ims_timeline)})",
        )
        local.stop()
    finally:
        volume_cache.MIN_CACHE_BYTES = floor
    app.timeline_manager.stop()
    print(flush=True)

    print("movie export", flush=True)
    encoder_ok, encoder_message = movie.encoder_available()
    if not gl_available:
        print("  skip movie checks — no OpenGL context", flush=True)
    else:
        with tempfile.TemporaryDirectory() as directory:
            spec = movie.MovieSpec(
                path=Path(directory) / "series.mov" if encoder_ok else Path(directory) / "series.gif",
                fps=4, start=0, stop=2, scale=1, timestamp=True, interval_s=2.0,
            )
            if not encoder_ok:
                print(f"  skip .mov — {encoder_message}; writing a GIF instead", flush=True)
            seen: list[int] = []
            before_export = timeseries.current_index(app.viewer, axis)
            written = movie.export_movie(app.viewer, spec, on_progress=lambda done, total: seen.append(done))
            check(written.exists() and written.stat().st_size > 500, f"{written.name} written ({written.stat().st_size} bytes)")
            check(seen == [1, 2, 3], f"progress was reported per frame ({seen})")
            check(
                timeseries.current_index(app.viewer, axis) == before_export,
                "the viewer was put back on the timepoint it started on",
            )

        # Cancelling has to leave nothing behind rather than a truncated file.
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "cancelled.mov"
            try:
                movie.export_movie(
                    app.viewer,
                    movie.MovieSpec(path=path, fps=4, start=0, stop=2, scale=1),
                    should_cancel=lambda: True,
                )
                check(False, "a cancelled export should raise")
            except movie.MovieCancelled:
                check(not path.exists(), "a cancelled export writes nothing")
    print(flush=True)

    print("toolbar actions", flush=True)
    visible = app.viewer.scale_bar.visible
    app.toolbar.toggle_scale_bar()
    check(app.viewer.scale_bar.visible != visible, "scale bar toggled")
    app.toolbar.toggle_scale_bar()
    # isHidden rather than isVisible: this viewer was built with show=False, so
    # every child widget reports itself invisible regardless of the dock state.
    app.toggle_metadata_panel()
    check(app._metadata_dock.isHidden(), "metadata dock hidden")
    app.toggle_metadata_panel()
    check(not app._metadata_dock.isHidden(), "metadata dock shown again")
    check("slide" in app.toolbar._buttons, "the slide export button is on the toolbar")
    check("movie" in app.toolbar._buttons, "and the movie export button")
    app.toolbar.toggle_play()
    check(app.timeseries_widget.playing, "the toolbar's play button starts playback")
    app.toolbar.toggle_play()
    check(not app.timeseries_widget.playing, "and stops it again")
    print(flush=True)

    print("the startup splash", flush=True)
    from microscopy_viewer import splash as splash_module

    check(splash_module.ANIMATION.is_file(), f"the animation ships with the package ({splash_module.ANIMATION.name})")
    banner = splash_module.start("Smoke test")
    check(banner is not None, "a splash is made when a QApplication exists")
    if banner is not None:
        check(banner.widget.isVisible(), "and it is on screen")
        check(banner._movie is not None, "Qt can play the animation")
        if banner._movie is not None:
            check(banner._movie.frameCount() > 1, f"which has {banner._movie.frameCount()} frames")
        # The whole point of pump(): the build never returns to the event loop,
        # so this is what makes the splash paint and say where it has got to.
        banner.pump("Building the Segmentation panel…")
        check(banner._message.text().startswith("Building"), "pump puts the step on the splash")
        banner.finish(None)
        check(not banner.widget.isVisible(), "and it closes when the window is ready")

    # The progress hook is the same one the splash is driven through, and a
    # viewer built without one must behave exactly as it always did.
    steps: list[str] = []
    quiet = MicroscopyViewer(show=False, progress=steps.append)
    check(steps and steps[0].startswith("Starting"), f"the build reports its steps ({len(steps)} of them)")
    check(any("panel" in step for step in steps), "including the panel being built")
    check(steps[-1] == "Ready", f"and says when it is done ({steps[-1]!r})")
    quiet.viewer.close()
    print(flush=True)

    print("PowerPoint slide export", flush=True)
    from microscopy_viewer import slides
    from microscopy_viewer.widgets.slide_dialog import SlideExportDialog

    dialog = SlideExportDialog(app.viewer, None)
    check(dialog._sample_table.rowCount() > 0, f"the dialog lists {dialog._sample_table.rowCount()} sample(s)")
    check(dialog._channel_table.rowCount() > 0, f"and {dialog._channel_table.rowCount()} channel(s)")
    check(bool(dialog.labels()), "channel labels are pre-filled from the file")
    # -- manual contrast ------------------------------------------------------
    from qtpy.QtCore import Qt

    from microscopy_viewer.widgets.slide_dialog import _parse_limits

    table = dialog._channel_table
    check(table.columnCount() == 5, f"the channel table carries Min and Max ({table.columnCount()} columns)")
    check(not dialog.contrast_limits(), "nothing is typed until manual mode is chosen")
    check(not (table.item(0, 3).flags() & Qt.ItemIsEditable), "and the limit cells start locked")

    dialog._contrast.setCurrentText(slides.CONTRAST_MANUAL)
    check(bool(table.item(0, 3).flags() & Qt.ItemIsEditable), "manual mode unlocks them")
    check(table.item(0, 4).text() != "", f"seeded from the displayed range ({table.item(0, 4).text()!r})")

    key = dialog._columns[0][0]
    table.item(0, 3).setText("10")
    table.item(0, 4).setText("2000")
    check(dialog.contrast_limits().get(key) == (10.0, 2000.0), "what is typed is what the export gets")

    # A backwards or half-typed pair falls back rather than raising: the picture
    # is then the displayed range, which is better than a modal complaint.
    table.item(0, 3).setText("2000")
    table.item(0, 4).setText("10")
    check(key not in dialog.contrast_limits(), "a backwards pair is ignored")
    check(_parse_limits("", "") is None and _parse_limits("abc", "5") is None, "so are empty and unparseable boxes")
    check(_parse_limits(" 3 ", "9.5") == (3.0, 9.5), "and a good pair reads as numbers")

    dialog._contrast.setCurrentText(slides.CONTRAST_AS_DISPLAYED)
    check(not dialog.contrast_limits(), "leaving manual mode drops the typed limits again")

    chosen = dialog.selected_samples()
    check(len(chosen) == dialog._sample_table.rowCount(), "every sample starts ticked")
    dialog.close()

    with tempfile.TemporaryDirectory() as directory:
        written = slides.export_slide(
            chosen, Path(directory) / "smoke_slide.pptx", title="Smoke test", max_pixels=200
        )
        check(
            written.exists() and written.stat().st_size > 10_000,
            f"a slide was written ({written.stat().st_size} bytes)",
        )

        from pptx import Presentation

        shapes = Presentation(str(written)).slides[0].shapes
        table = [shape.table for shape in shapes if shape.has_table][0]
        check(len(table.rows) == len(chosen) + 1, f"a header row plus one per sample ({len(table.rows)})")
        check(table.cell(0, len(table.columns) - 1).text == "Merge", "the last column is the merge")
        check(len([s for s in shapes if s.shape_type == 13]) > 0, "pictures were placed on the slide")
    print(flush=True)

    print("batch mode: a folder straight into the dialog", flush=True)
    batch = SlideExportDialog(None, SAMPLES)
    check(not batch.selected_samples(), "an empty dialog starts with nothing to export")
    before = batch._sample_table.rowCount()
    batch._load([SAMPLES])
    added = batch._sample_table.rowCount()
    check(added > before, f"the folder added {added - before} sample(s) without touching the viewer")
    check(batch._channel_table.rowCount() > 0, "and populated the channel table")
    check("slide" in batch._status.text(), f"the summary says how many slides ({batch._status.text()})")

    # A second pass over the same folder must not duplicate rows.
    batch._load([SAMPLES])
    check(batch._sample_table.rowCount() == added, "loading the same folder again adds nothing")

    # An edited label must survive more files being added.
    key = batch._columns[0][0]
    batch._channel_table.item(0, 1).setText("anti-CD31")
    batch._load([SAMPLES / "sample_zstack_imagej.tif"])
    check(batch.labels()[key] == "anti-CD31", "a typed channel label survives a further load")

    batch._rows_per_slide.setValue(2)
    with tempfile.TemporaryDirectory() as directory:
        # Closeups off: this check is about how the rows are split, and a folder
        # that happens to hold a nested acquisition would add slides to the count.
        deck = slides.export_slide(
            batch.selected_samples(), Path(directory) / "batch.pptx",
            max_pixels=120, rows_per_slide=2, contrast=slides.CONTRAST_AUTO,
            zoom_slides=False,
        )
        count = len(Presentation(str(deck)).slides)
        expected = -(-len(batch.selected_samples()) // 2)
        check(count == expected, f"{count} slide(s) at two rows each, as expected ({expected})")
    batch.close()
    print(flush=True)

    print("batch maximum projection", flush=True)
    from microscopy_viewer import projection as pj
    from microscopy_viewer.widgets.projection_dialog import ProjectionDialog

    check("mip" in app.toolbar._buttons, "the batch MIP button is on the toolbar")
    with tempfile.TemporaryDirectory() as directory:
        stacks = Path(directory) / "stacks"
        stacks.mkdir()
        for name in ("fish_1", "fish_2"):
            shutil.copy(SAMPLES / "sample_4d_2ch.ims", stacks / f"{name}.ims")
        out = Path(directory) / "MIP"

        layers_before = len(app.viewer.layers)
        dialog = ProjectionDialog(stacks)
        try:
            dialog._input_edit.setText(str(stacks))
            dialog.rescan()
            for _ in range(1500):
                QCoreApplication.processEvents()
                if dialog._channel_list.count():
                    break
                time.sleep(0.01)
            offered = [dialog._channel_list.item(row).text() for row in range(dialog._channel_list.count())]
            check(offered == ["GFP", "mCherry"], f"channel names were read from the files ({offered})")
            check(len(dialog._paths) == 2, f"both stacks were listed ({len(dialog._paths)})")

            # Tick one channel by name; the point of the list is picking stains,
            # not positions.
            dialog._output_edit.setText(str(out))
            dialog._all_channels.setChecked(False)
            for row in range(dialog._channel_list.count()):
                item = dialog._channel_list.item(row)
                item.setCheckState(Qt.Checked if item.text() == "mCherry" else Qt.Unchecked)
            check(dialog.chosen_channels() == ("mCherry",), "one channel ticked")

            dialog._format_box.setCurrentIndex(1)  # Imaris
            check(dialog.options().fmt == ".ims", f"the format combo yields a suffix ({dialog.options().fmt})")

            dialog.run()
            for _ in range(3000):
                QCoreApplication.processEvents()
                if dialog._worker is None and dialog.outcomes:
                    break
                time.sleep(0.01)
            written = sorted(path.name for path in out.iterdir())
            check(written == ["fish_1_MIP.ims", "fish_2_MIP.ims"], f"one output per stack ({written})")

            source = load_and_release(stacks / "fish_1.ims")
            result = load_and_release(out / "fish_1_MIP.ims")
            check(
                [spec.channel_name for spec in result] == ["mCherry"],
                f"carrying only the ticked channel ({[s.channel_name for s in result]})",
            )
            reference = next(spec for spec in source if spec.channel_name == "mCherry")
            expected = np.asarray(reference.data).max(axis=reference.axes.index("Z"))
            check(
                np.array_equal(np.asarray(result[0].data), expected),
                "and it really is the maximum over Z, pixel for pixel",
            )
            check(
                tuple(round(float(v), 4) for v in result[0].scale)[-2:] == (0.13, 0.13),
                f"with the pixel size intact ({tuple(round(float(v), 4) for v in result[0].scale)})",
            )
            check(
                len(app.viewer.layers) == layers_before,
                f"and nothing was added to the viewer ({len(app.viewer.layers)})",
            )

            # A second pass must not silently rewrite what is already there.
            dialog.run()
            for _ in range(2000):
                QCoreApplication.processEvents()
                if dialog._worker is None and dialog.outcomes:
                    break
                time.sleep(0.01)
            check(
                all(outcome.skipped for outcome in dialog.outcomes),
                "re-running skips outputs that already exist",
            )
            check("already existed" in dialog._status.text(), f"and says so ({dialog._status.text()[:60]!r})")
        finally:
            dialog.close()
            for path in list(stacks.glob("*.ims")) + list(out.glob("*.ims")):
                loader_release(path)
    print(flush=True)

    print("drag-and-drop path handling", flush=True)
    before = len(app.viewer.layers)
    app.open_paths([str(SAMPLES / "sample_zstack_imagej.tif")])
    check(len(app.viewer.layers) == before + 1, f"dropped file added one layer ({len(app.viewer.layers)})")

    from microscopy_viewer.loaders import expand_inputs, is_supported

    check(is_supported(SAMPLES / "sample_plain.tif"), "TIFF recognised as supported")
    check(not is_supported(SAMPLES / "nope.docx"), "unrelated file rejected")
    expanded = expand_inputs([SAMPLES])
    check(len(expanded) >= 5, f"dropping the folder expands to {len(expanded)} files")
    print(flush=True)

    print("panel manifest", flush=True)
    from microscopy_viewer.widgets.registry import iter_panels

    identifiers = [spec.identifier for spec in iter_panels()]
    check("metadata" in identifiers, f"metadata panel registered ({identifiers})")
    check("measurements" in identifiers, "measurements panel registered")
    check("intensity_comparison" in identifiers, "intensity comparison panel registered")
    check("timeseries" in identifiers, "time-series panel registered")
    check("atlas_registration" in identifiers, "atlas registration panel registered")
    check("segmentation" in identifiers, "segmentation panel registered")
    check("regions" in identifiers, "brain regions panel registered")
    check("experiment" in identifiers, "experiment setup panel registered")
    check("analysis" in identifiers, "analysis panel registered")
    for identifier in identifiers:
        check(identifier in app.panels, f"{identifier} built and tracked in app.panels")
        check(identifier in app.docks, f"{identifier} has a dock")
    check(app.metadata_widget is not None, "metadata_widget attribute still populated")
    check(app.measurements_widget is not None, "measurements_widget attribute still populated")
    check(app.intensity_widget is not None, "intensity_widget attribute populated")
    check(app.registration_widget is not None, "registration_widget attribute populated")
    check(app.segmentation_widget is not None, "segmentation_widget attribute populated")
    check(app.regions_widget is not None, "regions_widget attribute populated")
    check(app.experiment_widget is not None, "experiment_widget attribute populated")
    print(flush=True)

    print("atlas registration panel", flush=True)
    from microscopy_viewer import registration as rg

    panel = app.registration_widget
    panel.refresh_layers()
    assigned = panel.roles()
    check(bool(assigned), f"the role table lists {len(assigned)} layer(s)")
    check(
        all(role in rg.ROLES for role in assigned.values()),
        f"every layer got a role from the manifest ({sorted(set(assigned.values()))})",
    )
    # Nothing in the sample data is a nuclear channel, so nothing may be promoted
    # to driving a registration by accident.
    check(
        rg.ROLE_DRIVER not in assigned.values(),
        "no driver is guessed when no channel looks nuclear",
    )
    check(rg.guess_role("Confocal - dapi") == rg.ROLE_DRIVER, "a DAPI channel is guessed as the driver")
    check(rg.guess_role("Confocal - Actub") == rg.ROLE_LANDMARK, "acetylated tubulin is guessed as a landmark")
    check(rg.guess_role("Confocal - sv2") == rg.ROLE_CARRY, "an unrecognised channel stays a carry-along")

    # Refusals have to be quiet status text, not modal dialogs: an offscreen run
    # would hang on one, and so would a user who just wanted to look at the panel.
    panel.run()
    # Which refusal comes back depends on what is installed: without antspyx it
    # names the package, with it the first missing thing is the atlas itself.
    refusal = panel._status.text().lower()
    check(
        any(word in refusal for word in ("driver", "antspyx", "atlas reference")),
        f"running with nothing set up explains itself ({panel._status.text()[:60]}…)",
    )
    panel.export_regions()
    check("Nothing to export" in panel._status.text(), "exporting with no result is refused politely")

    first_layer = next(iter(assigned))
    volume, problem = panel._volume_from_layer(app.viewer.layers[first_layer], rg.ROLE_CARRY)
    if volume is not None:
        check(len(volume.spacing) == 3, f"a snapshot carries a 3-axis voxel size ({volume.spacing})")
        check(
            all(size > 0 for size in volume.spacing),
            "and no zero voxel size, which would collapse the stack in physical space",
        )
    else:
        check("3D" in problem, f"a non-3D layer is reported rather than snapshotted ({problem})")

    if rg.backend_available("ants"):
        check(panel._run_button.isEnabled(), "antspyx is installed, so Register is live")
    else:
        check(not panel._run_button.isEnabled(), "without antspyx, Register is disabled")
        check(not panel._backend_notice.isHidden(), "and the panel says what to install")
    print(flush=True)

    print("segmentation panel", flush=True)
    from microscopy_viewer import segmentation as sg

    panel = app.segmentation_widget
    panel.refresh_layers()
    check(panel._channel_box.count() > 0, f"the channel list holds {panel._channel_box.count()} layer(s)")
    check(
        panel._measure_box.itemData(0) == "",
        "the measure list offers the segmented channel as its first entry",
    )
    panel.refresh_device()
    check(bool(panel._device_label.text()), f"the device is shown before any run ({panel._device_label.text()})")

    settings = panel.settings()
    check(settings.mode in sg.MODES, f"the mode combo yields a real mode ({settings.mode})")
    check(settings.diameter_um == 0.0, "the diameter starts on automatic rather than a made-up number")

    layer = app.viewer.layers[str(panel._channel_box.currentData())]
    image, voxel, problem = panel._snapshot(layer)
    if image is not None:
        check(image.ndim in (2, 3), f"a snapshot is 2D or 3D, whatever the layer was ({image.shape})")
        check(len(voxel) == image.ndim and all(size > 0 for size in voxel), f"with a voxel size ({voxel})")
    else:
        check("2D or 3D" in problem, f"an unusable layer is reported rather than snapshotted ({problem})")

    # Deliberately not calling run(): with cellpose installed it would download
    # model weights and segment for minutes, which is not what a smoke check is.
    panel.export_objects()
    check("Nothing to export" in panel._status.text(), "exporting with no result is refused politely")

    if sg.backend_available("cellpose"):
        check(panel._run_button.isEnabled(), "cellpose is installed, so Segment is live")
    else:
        check(not panel._run_button.isEnabled(), "without cellpose, Segment is disabled")
        check(not panel._backend_notice.isHidden(), "and the panel says what to install")

    # Maximum-projection mode, through the panel but around cellpose: the stub
    # keeps this a smoke check rather than a minutes-long segmentation.
    modes = [panel._mode_box.itemText(i) for i in range(panel._mode_box.count())]
    check(sg.MODE_MIP in modes, f"the projection mode is offered ({modes})")
    panel._mode_box.setCurrentText(sg.MODE_MIP)
    check(panel.settings().is_projection, "and the panel builds a projecting settings object")
    check(not panel._stitch.isEnabled(), "the stitch threshold is greyed out, since it is unused")

    if image is not None and image.ndim == 3:
        flat = sg.max_projection(image)
        check(flat.shape == image.shape[1:], f"a snapshot flattens to 2D ({flat.shape})")
        stub = _StubSegmentation()
        sg.register_backend(stub)
        try:
            projected = sg.segment_volume(
                image, voxel,
                settings=sg.SegmentationSettings(backend="stub", mode=sg.MODE_MIP),
            )
            check(projected.projected, "a projected run says so on the result")
            check(projected.masks.ndim == 2, f"and returns 2D labels ({projected.masks.shape})")
            stats = sg.object_table(
                projected.masks, flat, projected.voxel_size_um[-2:]
            )
            check(bool(stats), f"the object table measures on the projection ({len(stats)} row(s))")
            headers = sg.object_headers(2)
            check(headers["volume_um3"] == "Area (µm²)", "2D results are labelled as areas")
        finally:
            sg._BACKENDS.pop("stub", None)
    print(flush=True)

    print("brain regions panel", flush=True)
    from microscopy_viewer import regions as reg

    regions_panel = app.regions_widget
    labels = np.zeros((8, 64, 64), dtype=np.int32)
    labels[:, 4:8, 4:8] = 1        # centroid near (6, 6)
    labels[:, 4:8, 40:44] = 2      # centroid near (6, 41)
    labels[:, 40:44, 40:44] = 3    # centroid near (41, 41)
    from microscopy_viewer.loaders.layer_spec import world_units

    app.viewer.add_labels(
        labels, name="smoke labels", scale=(1.0, 1.0, 1.0), **world_units(app.viewer, 3)
    )
    regions_panel.refresh_layers()
    check(regions_panel._labels_box.count() > 0, "the label map list found a Labels layer")

    region_layer = regions_panel.region_layer(create=True)
    check(region_layer is not None, "the region layer is created on demand")

    # A Shapes layer added without units defaults to dimensionless "pixel", and
    # one of those in the list makes napari drop units for the whole viewer —
    # which puts the scale bar back to reading pixels over a calibrated image.
    def _world_units():
        units = app.viewer.layers.extent.units
        return None if units is None else tuple(str(unit) for unit in units)

    check(
        tuple(str(unit) for unit in region_layer.units) == ("micrometer", "micrometer"),
        f"the region layer took the images' unit ({tuple(str(u) for u in region_layer.units)})",
    )
    check(
        _world_units() is not None and _world_units()[-1] == "micrometer",
        f"and the viewer's units survived it ({_world_units()})",
    )
    roi_probe = mm.new_roi_layer(app.viewer, app.viewer.layers[0])
    check(
        _world_units() is not None and _world_units()[-1] == "micrometer",
        f"a measurements ROI layer does not break them either ({_world_units()})",
    )
    app.viewer.layers.remove(roi_probe)
    region_layer.add_rectangles(np.array([[0.0, 0.0], [0.0, 20.0], [64.0, 20.0], [64.0, 0.0]]))
    region_layer.add_rectangles(np.array([[0.0, 20.0], [0.0, 64.0], [64.0, 64.0], [64.0, 20.0]]))
    regions_panel.refresh_regions()
    check(
        regions_panel._region_table.rowCount() == 2,
        f"both outlines reached the table ({regions_panel._region_table.rowCount()})",
    )
    regions_panel.suggest_names()
    names = [regions_panel._region_table.item(row, 0).text() for row in range(2)]
    check(names == list(reg.SUGGESTED_REGIONS[:2]), f"names were suggested ({names})")

    # Rename through the table and check it reaches the layer, not just the cell.
    regions_panel._region_table.item(0, 0).setText("cerebellum left")
    check(
        regions_panel.collect_regions()[0].name == "cerebellum left",
        "editing the table renames the outline on the layer",
    )

    index = regions_panel._labels_box.findData("smoke labels")
    regions_panel._labels_box.setCurrentIndex(index)
    regions_panel.count()
    for _ in range(600):
        QCoreApplication.processEvents()
        if regions_panel._worker is None and regions_panel._counts:
            break
        time.sleep(0.01)
    counted = {row.region: row.n_objects for row in regions_panel._counts}
    check(sum(counted.values()) == 3, f"every object was counted once ({counted})")
    check(counted.get("cerebellum left") == 1, f"one object on the left ({counted})")
    check(
        regions_panel._result_table.rowCount() == len(regions_panel._counts),
        "the results table matches the counts",
    )
    with tempfile.TemporaryDirectory() as directory:
        target = Path(directory) / "counts.xlsx"
        from microscopy_viewer.exports import export_sheets

        written = export_sheets(
            {
                "Regions": reg.counts_dataframe(regions_panel._counts),
                "Objects": reg.objects_dataframe(
                    regions_panel._stats, regions_panel._regions
                ),
            },
            target,
        )
        check(Path(written).stat().st_size > 4000, f"a two-sheet workbook was written ({written.name})")
    app.viewer.layers.remove("smoke labels")
    print(flush=True)

    print("experiment setup panel", flush=True)
    from microscopy_viewer import ims_store as store
    from microscopy_viewer.loaders import ims as ims_reader

    exp_panel = app.experiment_widget
    with tempfile.TemporaryDirectory() as directory:
        folder = Path(directory)
        shutil.copy(SAMPLES / "sample_4d_2ch.ims", folder / "fish_1.ims")
        shutil.copy(SAMPLES / "sample_4d_2ch.ims", folder / "fish_2.ims")

        before_scan = len(app.viewer.layers)
        exp_panel._folder_edit.setText(str(folder))
        exp_panel.scan()
        for _ in range(1500):
            QCoreApplication.processEvents()
            if exp_panel._worker is None and exp_panel._entries:
                break
            time.sleep(0.01)
        check(len(exp_panel._entries) == 2, f"both samples were found ({len(exp_panel._entries)})")
        check(exp_panel._grid.count() == 2, "and both have a tile in the grid")
        check(all(entry.readable for entry in exp_panel._entries), "both read cleanly")
        icons = [not exp_panel._grid.item(row).icon().isNull() for row in range(2)]
        check(all(icons), f"thumbnails were drawn ({icons})")
        check(
            exp_panel._entries[0].n_channels == 2,
            f"channels counted without opening the file ({exp_panel._entries[0].n_channels})",
        )
        check(
            len(app.viewer.layers) == before_scan,
            f"scanning added nothing to the viewer ({len(app.viewer.layers)} layers)",
        )

        # Outlines on screen are written into the files themselves.
        exp_panel._grid.selectAll()
        written = exp_panel.current_rois()
        check(len(written) == 2, f"the panel picked up the outlines on the canvas ({len(written)})")
        exp_panel.write_rois_to_selected()
        stored = store.load_rois(folder / "fish_1.ims")
        check(len(stored) == 2, f"both ROIs reached the .ims file ({len(stored)})")
        check(
            stored[0].name == "cerebellum left",
            f"with the names they were given ({[roi.name for roi in stored]})",
        )
        check(
            len(store.load_rois(folder / "fish_2.ims")) == 2,
            "and into every selected sample, not just the first",
        )

        # The tile text is refreshed so the file's new contents are visible.
        check(
            "ROI" in exp_panel._grid.item(0).text(),
            f"the tile says what the file now carries ({exp_panel._grid.item(0).text()!r})",
        )

        # Reading them back onto the canvas is the other half of the round trip.
        region_layer.data = []
        exp_panel.load_rois_from_sample()
        check(len(region_layer.data) == 2, "the stored ROIs came back onto the canvas")

        # A label map written into a file has to be readable back out of it, or
        # a batch that says it saved 3 889 objects is indistinguishable from one
        # that saved nothing.
        masks = np.zeros((8, 16, 16), dtype=np.int32)
        masks[2:5, 2:6, 2:6] = 1
        masks[5:7, 9:13, 9:13] = 2
        store.save_labels(folder / "fish_1.ims", "DAPI labels", masks, (0.5, 0.2, 0.2))
        check(
            store.list_labels(folder / "fish_1.ims") == ["DAPI labels"],
            "a label map is stored in the file",
        )
        exp_panel._grid.setCurrentRow(0)
        before_labels = len(app.viewer.layers)
        exp_panel.load_labels_from_sample()
        added = [
            layer for layer in app.viewer.layers
            if layer.name.endswith("DAPI labels")
        ]
        check(len(added) == 1, f"and comes back as a layer ({len(app.viewer.layers) - before_labels})")
        if added:
            check(int(added[0].data.max()) == 2, "with the objects that were written")
            check(
                tuple(float(v) for v in added[0].scale) == (0.5, 0.2, 0.2),
                f"at the voxel size it was stored with ({tuple(added[0].scale)})",
            )
            app.viewer.layers.remove(added[0])

        # The batch borrows the segmentation panel's settings rather than
        # duplicating them.
        from microscopy_viewer.widgets.batch_segmentation_widget import channel_spec

        batch = app.segmentation_widget.batch
        check(batch is not None, "the segmentation panel has a Batch tab")
        check(
            [app.segmentation_widget._tabs.tabText(i)
             for i in range(app.segmentation_widget._tabs.count())][-1] == "Batch",
            "beside Setup and Objects",
        )
        check(not hasattr(exp_panel, "run_batch"), "and the experiment panel no longer has one")
        borrowed = batch.batch_settings()
        check(
            borrowed.mode == app.segmentation_widget.settings().mode,
            "the batch takes its settings from the Setup tab",
        )
        check(batch.options().channel == "dapi", "segments DAPI by default")
        batch._channel_edit.setText("")
        check(batch.options() is None, "and refuses to run without a channel")
        batch._channel_edit.setText("dapi")
        check(channel_spec("2") == 2, "a bare number is a channel index")
        check(channel_spec("dapi") == "dapi", "anything else is a channel name")
        check(channel_spec("  ") is None, "and an empty box is no channel at all")
        exp_panel._grid.clearSelection()
        exp_panel._worker = object()
        batch.run()
        check("already running" in batch._status.text(), "a batch will not start during a scan")
        exp_panel._worker = None

        # The check that matters: opening a sample must put its stored regions
        # back by itself. Saving them and then having to remember a Load button
        # is how a saved annotation looks lost. A second viewer, because the test
        # is that *opening a file* does it, not that a method works.
        fresh = MicroscopyViewer(show=False)
        try:
            fresh.open_paths([folder / "fish_1.ims"])
            restored = [region.name for region in fresh.regions_widget.collect_regions()]
            check(
                restored == ["cerebellum left", "midbrain"],
                f"opening the sample restored its stored regions ({restored})",
            )
            check(
                reg.REGION_LAYER_NAME in fresh.viewer.layers,
                "and made the region layer to put them on",
            )
            check(
                "Loaded 2 region(s)" in fresh.regions_widget._status.text(),
                f"and said so ({fresh.regions_widget._status.text()[:60]!r})",
            )

            # Outlines already on the canvas are somebody's unsaved work and must
            # not be silently replaced.
            fresh.regions_widget._regions_source = None
            fresh.open_paths([folder / "fish_1.ims"])
            check(
                "left alone" in fresh.regions_widget._status.text(),
                f"existing outlines are not clobbered ({fresh.regions_widget._status.text()[:60]!r})",
            )
        finally:
            # Closing the viewer does not give the file back: the reader keeps its
            # handle for the life of the process, which on Windows is enough to
            # stop the temporary folder being deleted.
            fresh.viewer.close()
            ims_reader.release(folder / "fish_1.ims")

        # Stepping through a folder replaces the sample on screen rather than
        # piling samples up — that is the whole point of the panel. Its own
        # viewer, because replacing closes every file-backed layer, including the
        # ones the sections after this one still need.
        stepper = MicroscopyViewer(show=False)
        try:
            panel = stepper.experiment_widget
            panel._folder_edit.setText(str(folder))
            panel.scan()
            for _ in range(2000):
                QCoreApplication.processEvents()
                if panel._worker is None and panel._entries:
                    break
                time.sleep(0.01)

            def _images():
                from napari.layers import Image

                return [layer for layer in stepper.viewer.layers if isinstance(layer, Image)]

            def _show(row):
                panel._grid.clearSelection()
                panel._grid.item(row).setSelected(True)
                panel.open_selected()

            _show(0)
            first = len(_images())
            check(first == 2, f"opening a sample shows its two channels ({first})")
            _show(1)
            check(
                len(_images()) == first,
                f"opening the next one replaces it rather than adding ({len(_images())})",
            )
            check(
                "Replaced" in panel._status.text() and "fish_2" in panel._status.text(),
                f"and says which sample is up ({panel._status.text()!r})",
            )
            check(
                not ims_reader.is_open(folder / "fish_1.ims"),
                "the sample that went away released its file",
            )

            # What a sample stores comes up with it: fish_1 carries two outlines
            # and a label map, both written earlier in this section.
            def _labels():
                from napari.layers import Labels

                return [layer.name for layer in stepper.viewer.layers if isinstance(layer, Labels)]

            regions_panel_2 = stepper.regions_widget
            _show(0)
            check(
                len(regions_panel_2.collect_regions()) == 2,
                f"opening a sample loads its stored regions ({len(regions_panel_2.collect_regions())})",
            )
            check(
                _labels() == ["fish_1 — DAPI labels"],
                f"and its stored label map ({_labels()})",
            )
            check(
                stepper.viewer.layers[-1].name == reg.REGION_LAYER_NAME,
                "with the outlines drawn on top",
            )

            # Switching takes all of it away — the old outlines must not stay on
            # screen and block the next sample's own.
            fish_2 = folder / "fish_2.ims"
            store.save_rois(fish_2, [store.StoredRoi("hindbrain", np.array(
                [[0.0, 0.0], [0.0, 5.0], [5.0, 5.0], [5.0, 0.0]]))])
            _show(1)
            names = [region.name for region in regions_panel_2.collect_regions()]
            check(names == ["hindbrain"], f"the next sample shows its own regions, only ({names})")
            check(_labels() == [], f"and the previous sample's labels went with it ({_labels()})")
            check(len(_images()) == first, "and the swap still replaced the images")
            check(not regions_panel_2.has_unsaved_regions(), "freshly loaded outlines are not unsaved")

            # Unsaved drawing is never lost in silence: cancel keeps everything,
            # save writes it into the outgoing sample, discard drops it.
            layer = regions_panel_2.region_layer()
            layer.add_rectangles(np.array([[0.0, 0.0], [0.0, 20.0], [30.0, 20.0], [30.0, 0.0]]))
            regions_panel_2.refresh_regions()
            check(regions_panel_2.has_unsaved_regions(), "a new outline counts as unsaved")

            real_settle = panel._settle_regions
            try:
                panel._settle_regions = lambda _regions: None
                _show(0)
                check(
                    len(regions_panel_2.collect_regions()) == 2 and "fish_2" in str(regions_panel_2.sample_path()),
                    "cancelling keeps the sample and its drawing",
                )

                def _save(regions):
                    rois = [
                        store.StoredRoi(name=r.name or "extra", vertices_um=r.vertices_world)
                        for r in regions.collect_regions()
                    ]
                    return (fish_2, rois)

                panel._settle_regions = _save
                _show(0)
                check(
                    len(store.load_rois(fish_2)) == 2,
                    f"saving writes the drawing into the outgoing sample ({len(store.load_rois(fish_2))})",
                )
                check(
                    [region.name for region in regions_panel_2.collect_regions()]
                    == ["cerebellum left", "midbrain"],
                    "and the incoming sample still shows its own",
                )

                regions_panel_2.region_layer().add_rectangles(
                    np.array([[0.0, 0.0], [0.0, 9.0], [9.0, 9.0], [9.0, 0.0]])
                )
                panel._settle_regions = lambda _regions: ()
                _show(1)
                check(
                    len(store.load_rois(folder / "fish_1.ims")) == 2,
                    "discarding writes nothing",
                )
            finally:
                panel._settle_regions = real_settle

            # Selecting two shows exactly two.
            panel._grid.clearSelection()
            for row in (0, 1):
                panel._grid.item(row).setSelected(True)
            panel.open_selected()
            check(len(_images()) == 2 * first, f"two selected shows both ({len(_images())})")

            # The analysis panel reads the selection straight out of the files.
            analysis = stepper.analysis_widget
            check(analysis is not None, "analysis panel built")
            panel._grid.selectAll()
            analysis.run()
            for _ in range(3000):
                QCoreApplication.processEvents()
                if analysis._worker is None and analysis._outcomes:
                    break
                time.sleep(0.01)
            check(len(analysis._outcomes) == 2, f"both samples analysed ({len(analysis._outcomes)})")
            check(
                all(outcome.ok for outcome in analysis._outcomes),
                f"without errors ({[o.error for o in analysis._outcomes]})",
            )
            check(analysis._table.rowCount() == 2, "and listed in the panel")
            check(len(_images()) == 2 * first, "analysing did not disturb the open samples")
            try:
                np.asarray(_images()[0].data[0][0] if _images()[0].multiscale else _images()[0].data[0])
                readable = True
            except Exception:
                readable = False
            check(readable, "their files are still readable on screen")
            auto = analysis.last_report
            check(
                auto is not None and auto.parent == folder and any(auto.glob("*.png")),
                f"the run left a report folder with figures in the experiment folder ({auto})",
            )
            elsewhere = folder / "elsewhere"
            elsewhere.mkdir()
            saved = analysis.export(str(elsewhere))
            if saved is not None:
                import pandas as pd

                books = list(saved.glob("*.xlsx"))
                sheets = pd.read_excel(books[0], sheet_name=None) if books else {}
                check(
                    {"Region features", "Region intensities", "PCA matrix"} <= set(sheets),
                    f"the analysis workbook has its extra sheets ({sorted(sheets)})",
                )
            else:
                check(False, "the report was saved elsewhere too")
        finally:
            stepper.viewer.close()
            for name in ("fish_1.ims", "fish_2.ims"):
                ims_reader.release(folder / name)

        # A sample carrying nothing must not conjure an empty region layer.
        bare = MicroscopyViewer(show=False)
        try:
            bare.open_paths([SAMPLES / "sample_plain.tif"])
            check(
                bare.regions_widget.region_layer() is None,
                "a sample with no stored regions adds no region layer",
            )
        finally:
            bare.viewer.close()

    app.viewer.layers.remove(reg.REGION_LAYER_NAME)
    print(flush=True)

    print("window fits a laptop screen", flush=True)
    from qtpy.QtWidgets import QScrollArea

    smallest = app.viewer.window._qt_window.minimumSizeHint()
    # A 13" MacBook Air leaves about 1280 x 740 for the window. The panels used
    # to add up to far more, which left the window unresizable on one screen.
    check(
        smallest.width() <= 1280 and smallest.height() <= 720,
        f"the window can shrink to {smallest.width()} x {smallest.height()}",
    )
    check(
        all(isinstance(dock.widget(), QScrollArea) for dock in app.docks.values()),
        "every panel scrolls instead of holding the window open",
    )
    print(flush=True)

    print("files handed over by macOS (dropped on the app icon)", flush=True)
    from qtpy.QtCore import QEvent
    from qtpy.QtWidgets import QApplication

    from microscopy_viewer.dragdrop import catch_file_open_events

    class FileOpen:
        """Stands in for QFileOpenEvent, which PyQt5 will not construct."""

        def __init__(self, path):
            self._path = path

        def type(self):
            return QEvent.FileOpen

        def file(self):
            return self._path

    qapp = QApplication.instance()
    catcher = catch_file_open_events()
    received: list[list[str]] = []
    for path in ("/data/a.tif", "/data/b.ims", sys.argv[0]):
        catcher.eventFilter(qapp, FileOpen(path))
    QCoreApplication.processEvents()
    check(received == [], "held until the viewer is ready")
    catcher.attach(received.append)
    for _ in range(5):
        QCoreApplication.processEvents()
    check(
        received == [["/data/a.tif", "/data/b.ims"]],
        f"then opened together, without the launcher script ({received})",
    )
    qapp.removeEventFilter(catcher)
    catcher.deleteLater()
    print(flush=True)

    print("guided tour", flush=True)
    from qtpy.QtCore import QPoint, QRect

    from microscopy_viewer import onboarding as ob
    from microscopy_viewer.widgets.tour import TourOverlay

    real_state_file = ob.state_file
    with tempfile.TemporaryDirectory() as directory:
        # Never record anything in the real settings of whoever runs this.
        ob.state_file = lambda: Path(directory) / "onboarding.json"
        guided = MicroscopyViewer(show=True)
        try:
            guided.viewer.window._qt_window.resize(1400, 900)
            for _ in range(60):
                QCoreApplication.processEvents()
                time.sleep(0.01)
            check("tour" in guided.toolbar._buttons, "the toolbar has a Tour button")
            check(not ob.has_seen(), "a fresh profile has not seen the tour")
            tour = guided.maybe_start_tour()
            check(isinstance(tour, TourOverlay), "so it starts by itself")
            from microscopy_viewer.widgets.tour import _screen_rect

            missing = []
            off_screen = []
            for index, step in enumerate(ob.TOUR):
                tour.go(index)
                for _ in range(10):
                    QCoreApplication.processEvents()
                if step.target and (tour.target is None or not tour.target.isVisible()):
                    missing.append(step.target)
                bubble = tour._bubble
                placed = QRect(bubble.mapToGlobal(QPoint(0, 0)), bubble.size())
                if not _screen_rect(tour).contains(placed):
                    off_screen.append(index)
            check(missing == [], f"every step finds its control, on screen ({missing})")
            check(off_screen == [], f"every bubble is on the screen ({off_screen})")

            tour.go(2)
            for _ in range(10):
                QCoreApplication.processEvents()
            target = tour.target
            centre = tour.mapFromGlobal(target.mapToGlobal(target.rect().center()))
            check(not tour.mask().contains(centre), "the highlighted control can be clicked")
            check(tour.mask().contains(QPoint(2, tour.height() - 2)), "the rest is blocked")
            check(
                guided.docks["experiment"].isVisible(),
                "the step's panel was brought forward",
            )
            tour.go(8)
            tabs = guided.segmentation_widget._tabs
            check(tabs.tabText(tabs.currentIndex()) == "Batch", "and the step's tab chosen")

            outcomes = []
            tour.finished.connect(outcomes.append)
            tour._escape.activated.emit()
            QCoreApplication.processEvents()
            check(outcomes == [False], f"Esc stops the tour ({outcomes})")
            check(guided._tour is None and not tour.isVisible(), "and takes the overlay away")
            check(ob.has_seen(), "a stopped tour is not started again by itself")
            check(guided.maybe_start_tour() is None, "…and indeed is not")

            again = guided.start_tour()
            done = []
            again.finished.connect(done.append)
            for _ in range(len(ob.TOUR)):
                again.next()
            QCoreApplication.processEvents()
            check(done == [True], f"the Tour button runs it again, to the end ({done})")
            replaced = guided.start_tour(3)
            newer = guided.start_tour()
            check(guided._tour is newer and not replaced.isVisible(), "starting twice replaces the first")
            newer.end()
            newer.end()  # a second stop is harmless
        finally:
            ob.state_file = real_state_file
            guided.viewer.close()
    print(flush=True)

    print("ROI intensity comparison panel", flush=True)
    from qtpy.QtCore import Qt
    from qtpy.QtWidgets import QTableWidgetItem

    from microscopy_viewer import intensity as ix

    panel = app.intensity_widget
    panel.fill_from_layers()
    rows = panel.condition_rows()
    check(len(rows) >= 3, f"one condition per image layer ({len(rows)})")
    check(all(name and layer for name, layer in rows), f"every row is populated: {rows[:2]}")

    # Two conditions is the minimum the panel needs; keep the first two layers.
    while panel._condition_table.rowCount() > 2:
        panel._condition_table.selectRow(panel._condition_table.rowCount() - 1)
        panel.remove_condition()
    check(len(panel.condition_rows()) == 2, f"trimmed to two conditions ({len(panel.condition_rows())})")

    print("  -- projection layers --", flush=True)
    source_name = panel.condition_rows()[0][1]
    source = app.viewer.layers[source_name]
    check(source.ndim == 4, f"the source layer is a 4D stack ({source.ndim}D)")
    before = len(app.viewer.layers)

    panel._mode_box.setCurrentText("Maximum projection")
    panel.create_projection_layers()
    for _ in range(600):
        QCoreApplication.processEvents()
        if len(app.viewer.layers) > before:
            break
        time.sleep(0.01)

    mip_name = ix.projection_layer_name(source_name, "Maximum projection")
    check(mip_name in app.viewer.layers, f"a MIP layer was added ({mip_name})")
    mip = app.viewer.layers[mip_name]
    check(mip.ndim == 2, f"the MIP layer is 2D so shapes can be drawn on it ({mip.ndim}D)")
    check(tuple(mip.data.shape) == (128, 160), f"full XY resolution kept {tuple(mip.data.shape)}")
    check(tuple(mip.scale) == tuple(source.scale)[-2:], f"calibration carried over {tuple(mip.scale)}")
    check(mip.metadata.get("mv_projection") == "Maximum projection", "projection recorded on the layer")
    check(mip.metadata.get("mv_source_layer") == source_name, "source layer recorded")
    check(mip.metadata.get("mv_axes") == "YX", "axes updated to YX")

    # The projection must be the true max over Z at the displayed timepoint.
    step = int(app.viewer.dims.current_step[0])
    expected = np.asarray(source.data[0][step]).max(axis=0)
    check(np.array_equal(np.asarray(mip.data), expected), "pixels equal the max over Z")

    check(
        panel.condition_rows()[0][1] == mip_name,
        f"the condition now points at the MIP layer ({panel.condition_rows()[0][1]})",
    )
    check(panel._mode_box.currentText() == "Current slice", "mode reset now that the layer is flat")

    # Running it twice must update in place rather than pile up duplicates.
    count = len(app.viewer.layers)
    panel._mode_box.setCurrentText("Maximum projection")
    panel.create_projection_layers()
    for _ in range(400):
        QCoreApplication.processEvents()
        if panel._project_button.isEnabled():
            break
        time.sleep(0.01)
    check(len(app.viewer.layers) == count, f"re-projecting adds no duplicate layers ({len(app.viewer.layers)})")
    print(flush=True)

    roi_layer = panel.create_roi_layer()
    check(roi_layer is not None and roi_layer.name == "Comparison ROIs", "shared ROI layer created")
    check(panel.create_roi_layer() is roi_layer, "calling again reuses the same layer")

    # The conditions now point at the 2D MIP layers, so the ROI layer is 2D too.
    check(roi_layer.ndim == 2, f"the ROI layer matches the projected layers ({roi_layer.ndim}D)")
    signal_roi = np.array([[10, 10], [30, 10], [30, 30], [10, 30]], dtype=float)
    background_roi = np.array([[60, 60], [80, 60], [80, 80], [60, 80]], dtype=float)
    roi_layer.add_rectangles(signal_roi)
    roi_layer.add_rectangles(background_roi)
    mm.ensure_name_feature(roi_layer)
    panel.refresh_rois()
    names = mm.roi_names(roi_layer)
    check(len(names) == 2, f"two ROIs drawn ({names})")

    check(panel._roi_table.rowCount() == 2, f"the ROI table lists both ROIs ({panel._roi_table.rowCount()})")
    check(
        panel._roi_table.cellWidget(0, 1).currentText() == ix.ALL_CONDITIONS,
        "ROIs default to applying to every condition",
    )
    # Mark the second ROI as the background, shared across conditions.
    panel._roi_table.cellWidget(1, 2).setCurrentText("Background")

    conditions, rois, problems = panel._build_specs()
    check(len(conditions) == 2, f"two condition specs built ({len(conditions)})")
    check(len(rois) == 2, f"two ROI specs built ({len(rois)})")
    check(not problems, f"no setup problems reported ({problems})")
    check(sum(1 for roi in rois if roi.is_background) == 1, "exactly one ROI marked as background")
    check(
        conditions[0].data is not app.viewer.layers[conditions[0].layer_name].data
        or not app.viewer.layers[conditions[0].layer_name].multiscale,
        "multiscale layers are measured at full resolution, not a pyramid preview",
    )

    result = ix.run_comparison(conditions, rois, mode="Current slice")
    check(len(result.stats) == 4, f"two ROIs x two conditions ({len(result.stats)})")
    signal_stats = [s for s in result.stats if not s.is_background]
    check(all(s.n_pixels == 400 for s in signal_stats), f"20x20 ROI -> {[s.n_pixels for s in signal_stats]}")
    check(all(s.background_mean is not None for s in signal_stats), "background subtraction applied")

    panel._finish(result, problems)
    check(panel._stats_table.rowCount() == 4, f"statistics table filled ({panel._stats_table.rowCount()})")
    # One curve per ROI/condition, all ticked so the comparison is visible at once.
    check(panel._series_list.count() == 2, f"two curves listed ({panel._series_list.count()})")
    check(
        all(
            panel._series_list.item(row).checkState() == Qt.Checked
            for row in range(panel._series_list.count())
        ),
        "every curve starts ticked",
    )
    check(len(panel._selected_samples()[1]) == 2, "both curves feed the plot")
    check(panel._pair_a.count() == 2 and panel._pair_b.count() == 2, "curve pair combos populated")
    check(panel._pair_a.currentText() != panel._pair_b.currentText(), "two different curves preselected")
    check(len(panel._figure.axes) == 1, "histogram drawn onto the embedded canvas")
    check(len(panel._figure.axes[0].patches) > 0, "curves actually plotted")
    check(panel._separation.text() != "", f"separation metric shown: {panel._separation.text()[:60]}")

    # Unticking a curve drops it from the plot.
    panel._series_list.item(0).setCheckState(Qt.Unchecked)
    check(len(panel._selected_samples()[1]) == 1, "unticking removes a curve")
    panel._check_all(True)
    check(len(panel._selected_samples()[1]) == 2, "All restores them")

    panel._log_y.setChecked(True)
    panel._draw_histogram()
    check(panel._figure.axes[0].get_yscale() == "log", "log-y toggle applies to the plot")
    panel._log_y.setChecked(False)

    # The projection modes must not raise on a 4D layer.
    for mode in ix.PLANE_MODES:
        panel._mode_box.setCurrentText(mode)
        plane, description = ix.extract_plane(conditions[0], mode)
        check(plane.ndim == 2, f"{mode} yields a 2D plane {plane.shape} ({description})")
    panel._mode_box.setCurrentText("Current slice")

    print("  -- per-condition ROIs and normalisation --", flush=True)
    conditions_named = [name for name, _layer in panel.condition_rows()]
    # Give each condition its own signal ROI, under a shared comparison label.
    panel._roi_table.cellWidget(0, 2).setCurrentText("Signal")
    panel._roi_table.cellWidget(1, 2).setCurrentText("Signal")
    panel.assign_one_per_condition()
    assigned = panel._roi_assignments()
    applied = {values[0] for values in assigned.values()}
    check(applied == set(conditions_named), f"one ROI per condition ({applied})")
    labels = {values[2] for values in assigned.values()}
    check(labels == {"Sample 1"}, f"both share a comparison label ({labels})")

    _conditions, per_rois, _problems = panel._build_specs()
    check(all(roi.condition is not None for roi in per_rois), "specs carry the condition assignment")
    check(all(roi.label == "Sample 1" for roi in per_rois), "specs carry the shared label")
    for name in conditions_named:
        check(len(ix.rois_for(per_rois, name)) == 1, f"exactly one ROI applies to {name}")

    per_result = ix.run_comparison(_conditions, per_rois, mode="Current slice")
    check(len(per_result.stats) == 2, f"one measurement per condition ({len(per_result.stats)})")
    check(
        {s.roi_name for s in per_result.stats} == {"Sample 1"},
        "both rows share the label so they can be overlaid",
    )
    check(
        len({s.roi_source for s in per_result.stats}) == 2,
        "but they came from two different shapes",
    )
    panel._finish(per_result, [])
    # The point of the exercise: two ROIs drawn on two different conditions must
    # still show up as two overlaid curves.
    check(panel._series_list.count() == 2, f"both per-condition ROIs listed ({panel._series_list.count()})")
    check(len(panel._selected_samples()[1]) == 2, "and both are plotted together")
    check(len(panel._figure.axes[0].patches) > 0, "histogram drawn for per-condition ROIs")
    check(panel._pair_a.count() == 2, "either can be picked for the AUC")
    panel._update_separation()
    check("vs" in panel._separation.text(), f"separation computed across them: {panel._separation.text()[:50]}")

    # Restore a shared background so normalisation has something to work with.
    panel._roi_table.cellWidget(0, 1).setCurrentText(ix.ALL_CONDITIONS)
    panel._roi_table.cellWidget(1, 1).setCurrentText(ix.ALL_CONDITIONS)
    panel._roi_table.cellWidget(1, 2).setCurrentText("Background")
    # Give the background its own label; sharing one with a signal ROI would make
    # both write to the same (label, condition) sample key.
    panel._roi_table.setItem(1, 3, QTableWidgetItem("Background"))
    panel._normalization_box.setCurrentText("Divide by background")
    _conditions, norm_rois, _problems = panel._build_specs()
    norm_result = ix.run_comparison(
        _conditions, norm_rois, mode="Current slice", normalization="Divide by background"
    )
    signal_rows = [s for s in norm_result.stats if not s.is_background]
    check(all(s.normalized_mean is not None for s in signal_rows), "normalised means computed")
    check(all(s.normalization == "Divide by background" for s in signal_rows), "choice recorded per row")
    panel._finish(norm_result, [])
    check(
        panel._figure.axes[0].get_xlabel() == ix.normalized_axis_label("Divide by background"),
        f"the axis label reflects the normalisation ({panel._figure.axes[0].get_xlabel()!r})",
    )
    panel._normalization_box.setCurrentText("None")
    print(flush=True)

    with tempfile.TemporaryDirectory() as directory:
        csv_path = Path(directory) / "stats.csv"
        ix.stats_dataframe(norm_result.stats).to_csv(csv_path, index=False)
        header = csv_path.read_text(encoding="utf-8").splitlines()[0]
        check("Normalised mean" in header, f"CSV carries the normalised column: {header[:70]}")
        ix.stats_dataframe(result.stats).to_csv(csv_path, index=False)
        check(csv_path.stat().st_size > 0, f"stats CSV written ({csv_path.stat().st_size} bytes)")
        png_path = Path(directory) / "hist.png"
        panel._figure.savefig(png_path, dpi=100)
        check(png_path.stat().st_size > 1000, f"histogram PNG written ({png_path.stat().st_size} bytes)")
    print(flush=True)

    print("threaded measurement keeps the UI responsive", flush=True)
    panel._result = None
    panel.measure()
    check(not panel._measure_button.isEnabled(), "the Measure button is disabled while running")
    for _ in range(600):  # up to ~6 s
        QCoreApplication.processEvents()
        if panel._result is not None and panel._worker is None:
            break
        time.sleep(0.01)
    check(panel._result is not None, "the worker delivered a result back to the main thread")
    check(panel._worker is None, "the worker was cleared when it finished")
    check(panel._measure_button.isEnabled(), "the Measure button is re-enabled afterwards")
    check(panel._stats_table.rowCount() == 4, f"table refreshed from the worker ({panel._stats_table.rowCount()})")
    print(flush=True)

    app.viewer.close()

    if _failures:
        print(f"{len(_failures)} check(s) failed:")
        for failure in _failures:
            print(f"  - {failure}")
        return 1
    print("all GUI checks passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
