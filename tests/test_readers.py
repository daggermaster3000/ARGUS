"""Reader, measurement and export checks that need no display.

Run with::

    python tests/test_readers.py

(Also works under pytest, but the plain-script form is what the README documents
so it can be run in the same environment as the viewer.)
"""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

import numpy as np

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from microscopy_viewer.exports import export_measurements, measurements_dataframe  # noqa: E402
from microscopy_viewer.loaders import load_path  # noqa: E402
from microscopy_viewer.measurements import Measurement, measure_layer  # noqa: E402
from microscopy_viewer.utils import MICRON, MICRON_SQ  # noqa: E402

SAMPLES = Path(__file__).resolve().parent / "sample_data"

_failures: list[str] = []


def check(condition: bool, message: str) -> None:
    if condition:
        print(f"  ok   {message}")
    else:
        print(f"  FAIL {message}")
        _failures.append(message)


def close(actual, expected, tol=1e-6) -> bool:
    return actual is not None and abs(float(actual) - float(expected)) <= tol * max(1.0, abs(expected))


# ---------------------------------------------------------------------------


def test_ims_4d() -> None:
    print("sample_4d_2ch.ims (TZYX, 2 channels, 2 resolution levels)")
    specs = load_path(SAMPLES / "sample_4d_2ch.ims")
    check(len(specs) == 2, f"two layers returned (got {len(specs)})")
    spec = specs[0]
    check(spec.axes == "TZYX", f"axes are TZYX (got {spec.axes})")
    check(spec.multiscale, "multiscale pyramid detected")
    check(len(spec.data) == 2, f"two pyramid levels (got {len(spec.data)})")
    check(tuple(spec.data[0].shape) == (3, 8, 128, 160), f"level 0 shape {spec.data[0].shape}")
    check(tuple(spec.data[1].shape) == (3, 8, 64, 80), f"level 1 shape {spec.data[1].shape}")

    meta = spec.metadata
    check(close(meta.pixel_size_x_um, 0.13), f"pixel size X = {meta.pixel_size_x_um}")
    check(close(meta.pixel_size_y_um, 0.13), f"pixel size Y = {meta.pixel_size_y_um}")
    check(close(meta.z_step_um, 0.5), f"z step = {meta.z_step_um}")
    check(close(meta.time_interval_s, 30.0), f"time interval = {meta.time_interval_s}")
    check(meta.channel_names == ["GFP", "mCherry"], f"channel names {meta.channel_names}")
    check(close(meta.numerical_aperture, 1.40), f"NA = {meta.numerical_aperture}")
    check("63" in meta.objective, f"objective = {meta.objective!r}")
    check(meta.acquisition_date.startswith("2026-07-20"), f"date = {meta.acquisition_date!r}")
    check(meta.dimensionality.startswith("XYZT"), f"dimensionality = {meta.dimensionality!r}")

    gfp, mcherry = meta.channel(0), meta.channel(1)
    check(close(gfp.excitation_nm, 488), f"GFP excitation = {gfp.excitation_nm}")
    check(close(gfp.emission_nm, 509), f"GFP emission = {gfp.emission_nm}")
    check(close(gfp.laser_power, 3.2), f"GFP laser power = {gfp.laser_power}")
    check(close(gfp.exposure_ms, 120.0), f"GFP exposure = {gfp.exposure_ms} ms")
    check(close(mcherry.excitation_nm, 561), f"mCherry excitation = {mcherry.excitation_nm}")
    check(gfp.color == (0.0, 1.0, 0.0), f"GFP colour {gfp.color}")
    check(spec.scale == (30.0, 0.5, 0.13, 0.13), f"layer scale {spec.scale}")

    plane = np.asarray(spec.data[0][0, 0])
    check(plane.shape == (128, 160) and plane.max() > 0, "pixel data reads and is non-empty")


def test_ims_2d() -> None:
    print("sample_2d.ims (single plane, single channel)")
    specs = load_path(SAMPLES / "sample_2d.ims")
    check(len(specs) == 1, f"one layer returned (got {len(specs)})")
    spec = specs[0]
    check(spec.axes == "YX", f"singleton T and Z dropped, axes = {spec.axes}")
    check(len(spec.scale) == 2, f"scale has two entries {spec.scale}")
    shape = spec.data[0].shape if spec.multiscale else spec.data.shape
    check(tuple(shape) == (96, 128), f"shape {tuple(shape)}")


def test_brightfield_is_gray() -> None:
    print("sample_brightfield.ims (brightfield channel stored with a green colour)")
    specs = load_path(SAMPLES / "sample_brightfield.ims")
    check(len(specs) == 2, f"two layers returned (got {len(specs)})")
    brightfield, fluorescence = specs

    check(brightfield.channel_name == "Brightfield", f"channel 0 = {brightfield.channel_name!r}")
    check(brightfield.colormap == "gray", f"brightfield colormap = {brightfield.colormap!r}")
    check(
        brightfield.color is None,
        f"the file's green colour is discarded for brightfield (got {brightfield.color})",
    )

    # Without the brightfield rule, channel 0 would take the first cycle entry.
    check(fluorescence.channel_name == "GFP", f"channel 1 = {fluorescence.channel_name!r}")
    check(fluorescence.color == (0.0, 1.0, 0.0), f"GFP keeps its file colour {fluorescence.color}")


def test_brightfield_name_matching() -> None:
    print("brightfield channel-name matching")
    from microscopy_viewer.loaders.layer_spec import channel_appearance, is_brightfield

    for name in (
        "Brightfield", "bright field", "Bright-Field", "BF", "TL-BF", "TL",
        "DIC", "Phase", "PhC", "Transmitted", "Transmission", "T-PMT", "ESID", "Trans",
    ):
        check(is_brightfield(name), f"{name!r} recognised as brightfield")

    for name in (
        "GFP", "EGFP", "BFP", "mCherry", "DAPI", "Alexa 594", "Cy5",
        "Phalloidin", "TRITC", "Channel 0", "", None,
    ):
        check(not is_brightfield(name), f"{name!r} left as fluorescence")

    colormap, color, blending = channel_appearance("Brightfield", (0.0, 1.0, 0.0), 0, 3)
    check(colormap == "gray", f"brightfield -> {colormap!r}")
    check(color is None, f"stored colour dropped (got {color})")
    check(blending == "additive", f"blending unchanged for multichannel ({blending})")

    colormap, color, _blending = channel_appearance("GFP", None, 0, 3)
    check(colormap == "green", f"first fluorescence channel still cycles to {colormap!r}")
    check(color is None, "no colour invented when the file had none")


def test_pyramid_padding_without_imagesize_attrs() -> None:
    """Real Imaris files often omit ImageSize*, leaving chunk padding on coarse levels."""
    print("pyramid padding when ImageSize attrs are missing")
    import h5py
    import numpy as np

    from microscopy_viewer.loaders.ims import _real_size

    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "nosize.ims"
        with h5py.File(path, "w") as handle:
            # Stored shapes are padded up to the chunk size; the real image is
            # 2040 wide, exactly as in the deconvolved files from the microscope.
            level0 = handle.create_dataset("l0", shape=(24, 2048, 2048), dtype="u2")
            level1 = handle.create_dataset("l1", shape=(24, 1024, 1024), dtype="u2")
            image_info = handle.create_group("Image")
            for key, value in (("X", 2040), ("Y", 2040), ("Z", 17)):
                image_info.attrs.create(key, np.frombuffer(str(value).encode(), dtype="S1"))

            # Level 0 falls back to /DataSetInfo/Image and is cropped correctly.
            real0 = tuple(_real_size(level0, image_info, axis) for axis in "ZYX")
            check(real0 == (17, 2040, 2040), f"level 0 cropped to {real0}")

            # Level 1 must not reuse level 0's size, nor keep its own padding.
            reference = tuple(zip(real0, (24, 2048, 2048)))
            naive = tuple(_real_size(level1, image_info, axis) for axis in "ZYX")
            check(naive == (17, 1024, 1024), f"without a reference the padding survives {naive}")
            # Y and X halve; Z is stored at 24 on both levels, so it stays 17.
            derived = tuple(
                _real_size(level1, image_info, axis, reference[i]) for i, axis in enumerate("ZYX")
            )
            check(derived == (17, 1020, 1020), f"with a reference level 1 becomes {derived}")


def test_ome_tiff() -> None:
    print("sample_tzcyx.ome.tif (OME-TIFF, TZCYX, 2 channels)")
    specs = load_path(SAMPLES / "sample_tzcyx.ome.tif")
    check(len(specs) == 2, f"two layers returned (got {len(specs)})")
    spec = specs[0]
    check(spec.axes == "TZYX", f"channel axis split out, axes = {spec.axes}")
    shape = spec.data[0].shape if spec.multiscale else spec.data.shape
    check(tuple(shape) == (4, 6, 96, 128), f"shape {tuple(shape)}")

    meta = spec.metadata
    check(meta.file_format == "OME-TIFF", f"format = {meta.file_format}")
    check(close(meta.pixel_size_x_um, 0.108), f"pixel size = {meta.pixel_size_x_um}")
    check(close(meta.z_step_um, 0.35), f"z step = {meta.z_step_um}")
    check(close(meta.time_interval_s, 15.0), f"time interval = {meta.time_interval_s}")
    check(meta.channel_names == ["EGFP", "Alexa 594"], f"channel names {meta.channel_names}")
    check(close(meta.channel(0).exposure_ms, 50.0), f"exposure = {meta.channel(0).exposure_ms} ms")
    check(meta.acquisition_date.startswith("2026-07-21"), f"date = {meta.acquisition_date!r}")


def test_imagej_tiff() -> None:
    print("sample_zstack_imagej.tif (ImageJ Z stack)")
    specs = load_path(SAMPLES / "sample_zstack_imagej.tif")
    check(len(specs) == 1, f"one layer returned (got {len(specs)})")
    meta = specs[0].metadata
    check(specs[0].axes == "ZYX", f"axes = {specs[0].axes}")
    check(close(meta.z_step_um, 0.8), f"z step from ImageJ spacing = {meta.z_step_um}")
    check(close(meta.pixel_size_x_um, 0.2, tol=1e-4), f"pixel size from tags = {meta.pixel_size_x_um}")
    check(meta.is_calibrated, "reported as calibrated")


def test_plain_tiff() -> None:
    print("sample_plain.tif (no calibration)")
    specs = load_path(SAMPLES / "sample_plain.tif")
    meta = specs[0].metadata
    check(specs[0].axes == "YX", f"axes = {specs[0].axes}")
    check(not meta.is_calibrated, "reported as uncalibrated")
    check(specs[0].scale == (1.0, 1.0), f"scale falls back to 1 px {specs[0].scale}")


def test_ome_zarr() -> None:
    store = SAMPLES / "sample.zarr"
    if not store.exists():
        print("sample.zarr — skipped (not generated)")
        return
    print("sample.zarr (OME-Zarr, CZYX, 2 channels, 2 levels)")
    specs = load_path(store)
    check(len(specs) == 2, f"two layers returned (got {len(specs)})")
    spec = specs[0]
    check(spec.axes == "ZYX", f"channel axis split out, axes = {spec.axes}")
    check(spec.multiscale and len(spec.data) == 2, "two pyramid levels")
    meta = spec.metadata
    check(close(meta.pixel_size_x_um, 0.15), f"pixel size = {meta.pixel_size_x_um}")
    check(close(meta.z_step_um, 0.4), f"z step = {meta.z_step_um}")
    check(meta.channel_names == ["GFP", "RFP"], f"channel names {meta.channel_names}")


def test_ome_zarr_plate() -> None:
    store = SAMPLES / "sample_plate.zarr"
    if not store.exists():
        print("sample_plate.zarr — skipped (not generated)")
        return
    print("sample_plate.zarr (HCS plate, 1x2 wells, 2 fields, 2 channels)")
    specs = load_path(store)
    check(len(specs) == 2, f"one layer per channel (got {len(specs)})")
    spec = specs[0]
    meta = spec.metadata
    check(meta.file_format == "OME-Zarr (HCS plate)", f"format = {meta.file_format}")
    check(spec.axes == "YX", f"channel axis split out, axes = {spec.axes}")
    # 1 row x 2 columns of wells, each holding 2 fields side by side: 32 x (48*4).
    check(spec.data[0].shape == (32, 192), f"mosaic shape = {spec.data[0].shape}")
    check(spec.multiscale and len(spec.data) == 2, "two mosaic levels")
    check(spec.data[1].shape == (16, 96), f"level 1 mosaic shape = {spec.data[1].shape}")
    check(meta.channel_names == ["GFP", "mCherry"], f"channel names {meta.channel_names}")
    check(close(meta.pixel_size_x_um, 0.25), f"pixel size = {meta.pixel_size_x_um}")
    check(
        meta.extra.get("Wells assembled") == "2 (A/1 – A/2)",
        f"wells and plate region recorded = {meta.extra}",
    )
    tile = np.asarray(spec.data[0][:, :48])
    check(tile.max() > 0, "first field carries data")


def test_roi_layer_text_matches_the_layer_ndim() -> None:
    """The ROI label offset must have one entry per axis, not always two.

    napari indexes the text translation by displayed axis, so a two-entry offset
    on a 3D or 4D layer raises ``IndexError: index 2 is out of bounds`` as soon as
    a shape is drawn — the layer only has to be one plane deep to hit it.
    """
    print("ROI layers carry a text offset shaped like the image")
    import microscopy_viewer.measurements as mm

    class _Image:
        def __init__(self, ndim):
            self.ndim = ndim
            self.name = f"{ndim}D image"
            self.scale = tuple(1.0 for _ in range(ndim))
            self.metadata = {"mv_axes": "TZYX"[-ndim:], "mv_channel_name": "GFP"}

    class _Viewer:
        def __init__(self):
            self.captured = {}

        def add_shapes(self, **kwargs):
            self.captured = kwargs

            class _Layer:
                mode = ""
                name = kwargs["name"]

            return _Layer()

    for ndim in (2, 3, 4):
        viewer = _Viewer()
        mm.new_roi_layer(viewer, _Image(ndim))
        translation = viewer.captured["text"]["translation"]
        check(
            len(translation) == ndim,
            f"{ndim}D image -> {ndim}-entry translation (got {translation})",
        )
        check(
            list(translation[-2:]) == [-6.0, 0.0],
            f"and the label is still nudged in Y/X ({translation})",
        )


def test_ome_zarr_acquisitions_are_channels() -> None:
    """A 4i plate's cycles are further channels, not further places.

    The images of a well are listed the same way whether they are fields of view
    or staining cycles; only the ``acquisition`` id tells them apart. Tiling the
    cycles the way fields are tiled would lay the same cells out side by side and
    inflate the mosaic, so this pins the distinction down.
    """
    store = SAMPLES / "sample_fractal_plate.zarr"
    if not store.exists():
        print("sample_fractal_plate.zarr — skipped (not generated)")
        return
    print("sample_fractal_plate.zarr (2x2 wells, 2 cycles, 2 channels)")
    specs = load_path(store)

    check(len(specs) == 4, f"two cycles of two channels (got {len(specs)})")
    # 2 rows x 2 columns of wells, one field each: (32*2) x (48*2). A cycle tiled
    # as if it were a second field would double one of those.
    check(
        all(spec.data[0].shape == (64, 96) for spec in specs),
        f"every cycle covers the plate once: {[spec.data[0].shape for spec in specs]}",
    )
    check(
        [spec.name.split(" :: ", 1)[1] for spec in specs]
        == ["cycle 1 :: Ab1_DAPI", "cycle 1 :: Green488-x1",
            "cycle 2 :: Ab2_DAPI", "cycle 2 :: Green488-x2"],
        f"the cycle is in the layer name: {[spec.name for spec in specs]}",
    )
    check(
        specs[0].metadata.extra.get("Fields per well") == "1",
        f"one field per well, not seven: {specs[0].metadata.extra.get('Fields per well')}",
    )
    check(
        specs[0].metadata.extra.get("Acquisitions") == "2",
        f"the acquisition count is recorded: {specs[0].metadata.extra.get('Acquisitions')}",
    )

    well = load_path(store / "B" / "02")
    check(len(well) == 4, f"the well opens its 2 cycles x 2 channels (got {len(well)})")
    check(
        well[0].name == "B/02 :: cycle 1 :: Ab1_DAPI",
        f"a well layer names its cycle rather than the omero placeholder: {well[0].name}",
    )


def test_ome_zarr_sparse_plate_is_trimmed() -> None:
    """A 4x12 layout holding two wells assembles those two, not the whole plate."""
    store = SAMPLES / "sample_plate_sparse.zarr"
    if not store.exists():
        print("sample_plate_sparse.zarr — skipped (not generated)")
        return
    print("sample_plate_sparse.zarr (4x12 declared, 2 wells present)")
    specs = load_path(store)
    spec = specs[0]
    check(spec.data[0].shape == (32, 192), f"mosaic covers the acquired wells only {spec.data[0].shape}")
    extra = spec.metadata.extra
    check(extra.get("Plate layout") == "1 rows x 2 columns", f"layout = {extra.get('Plate layout')}")
    check(
        extra.get("Wells assembled") == "2 (A/1 – A/2)",
        f"the assembled region is named: {extra.get('Wells assembled')}",
    )


def test_ome_zarr_well_and_descent() -> None:
    """A well, a field and a plain row folder are all openable on their own."""
    store = SAMPLES / "sample_plate.zarr"
    if not store.exists():
        print("sample_plate.zarr — skipped (not generated)")
        return
    print("sample_plate.zarr subgroups (well, field, row)")

    well = load_path(store / "A" / "1")
    check(len(well) == 4, f"well opens its 2 fields x 2 channels (got {len(well)})")
    check(
        well[0].name == "A/1 :: field 0 :: GFP",
        f"well layer named by well and field, not by the store: {well[0].name}",
    )
    check(well[0].data[0].shape == (32, 48), f"field shape = {well[0].data[0].shape}")

    field = load_path(store / "A" / "1" / "0")
    check(len(field) == 2, f"a single field opens as its channels (got {len(field)})")
    check(field[0].name == "sample_plate :: A/1/0 :: GFP", f"field name {field[0].name}")

    # A row folder carries no NGFF metadata at all — the reader descends into it.
    row = load_path(store / "A")
    check(len(row) == 8, f"row folder descends to 2 wells x 2 fields x 2 channels (got {len(row)})")


def test_bioformats2raw_series() -> None:
    store = SAMPLES / "sample_series.zarr"
    if not store.exists():
        print("sample_series.zarr — skipped (not generated)")
        return
    print("sample_series.zarr (bioformats2raw layout, 2 series)")
    specs = load_path(store)
    check(len(specs) == 2, f"one layer per series (got {len(specs)})")
    check(
        [spec.name for spec in specs]
        == ["sample_series :: Well 0", "sample_series :: Well 1"],
        f"series named from METADATA.ome.xml: {[spec.name for spec in specs]}",
    )
    spec = specs[0]
    check(spec.axes == "YX", f"single channel axis split out, axes = {spec.axes}")
    check(spec.data[0].shape == (40, 56), f"series shape = {spec.data[0].shape}")


# ---------------------------------------------------------------------------


class _FakeShapes:
    """Minimal stand-in for a napari Shapes layer, so measuring needs no GUI."""

    def __init__(self, data, shape_type, scale, metadata):
        self.data = data
        self.shape_type = shape_type
        self.scale = scale
        self.metadata = metadata
        self.name = "ROIs"
        self.features = {}


def test_measurements() -> None:
    print("measurements on calibrated shapes")
    from microscopy_viewer.metadata import AcquisitionMetadata

    meta = AcquisitionMetadata(image_name="unit", pixel_size_x_um=0.5, pixel_size_y_um=0.5)
    # A 10x20 px rectangle and a 30 px horizontal line, at 0.5 µm/px.
    rectangle = np.array([[0, 0], [10, 0], [10, 20], [0, 20]], dtype=float)
    line = np.array([[5, 0], [5, 30]], dtype=float)
    layer = _FakeShapes(
        [rectangle, line],
        ["rectangle", "line"],
        (0.5, 0.5),
        {"mv_metadata": meta, "mv_image_name": "unit", "mv_channel_name": "GFP", "mv_axes": "YX"},
    )

    rows = measure_layer(layer)
    by_key = {(m.roi_name, m.measurement_type): m for m in rows}

    area = by_key[("ROI 1", "Area")]
    check(close(area.value, 50.0), f"rectangle area = {area.value} (expected 5 µm x 10 µm = 50)")
    check(area.unit == MICRON_SQ, f"area unit = {area.unit}")
    perimeter = by_key[("ROI 1", "Perimeter")]
    check(close(perimeter.value, 30.0), f"rectangle perimeter = {perimeter.value} (expected 30)")

    length = by_key[("ROI 2", "Length")]
    check(close(length.value, 15.0), f"line length = {length.value} (expected 30 px x 0.5 = 15)")
    check(length.unit == MICRON, f"length unit = {length.unit}")
    check(all(m.calibrated for m in rows), "all rows flagged calibrated")
    check(all(m.channel == "GFP" for m in rows), "channel recorded on every row")


def test_ellipse_and_uncalibrated() -> None:
    print("ellipse area and the uncalibrated fallback")
    from microscopy_viewer.metadata import AcquisitionMetadata

    # Bounding box 20 x 40 px, so semi-axes are 10 and 20 px.
    ellipse = np.array([[0, 0], [20, 0], [20, 40], [0, 40]], dtype=float)
    layer = _FakeShapes(
        [ellipse], ["ellipse"], (1.0, 1.0), {"mv_metadata": AcquisitionMetadata(), "mv_axes": "YX"}
    )
    rows = {m.measurement_type: m for m in measure_layer(layer)}
    check(close(rows["Area"].value, np.pi * 10 * 20), f"ellipse area = {rows['Area'].value}")
    check(close(rows["Major axis"].value, 40.0), f"major axis = {rows['Major axis'].value}")
    check(close(rows["Minor axis"].value, 20.0), f"minor axis = {rows['Minor axis'].value}")
    check(rows["Area"].unit == "px²", f"uncalibrated area unit = {rows['Area'].unit}")
    check(not rows["Area"].calibrated, "flagged as uncalibrated")


def test_3d_line_uses_z() -> None:
    print("3D line length includes the Z component")
    from microscopy_viewer.metadata import AcquisitionMetadata

    meta = AcquisitionMetadata(pixel_size_x_um=1.0, pixel_size_y_um=1.0, z_step_um=2.0)
    # From (z=0, y=0, x=0) to (z=3, y=4, x=0): 6 µm in Z, 4 µm in Y -> hypotenuse.
    line = np.array([[0, 0, 0], [3, 4, 0]], dtype=float)
    layer = _FakeShapes([line], ["line"], (2.0, 1.0, 1.0), {"mv_metadata": meta, "mv_axes": "ZYX"})
    rows = {m.measurement_type: m for m in measure_layer(layer)}
    check(close(rows["Length"].value, np.hypot(6.0, 4.0)), f"3D length = {rows['Length'].value}")


def test_3d_level_choice() -> None:
    print("3D pyramid level selection")
    from microscopy_viewer.rendering import choose_level, displayed_voxels

    class _Level:
        def __init__(self, shape):
            self.shape = shape

    # Leading (time) axes are sliced away before rendering and must not count.
    check(displayed_voxels((10, 4, 100, 200)) == 4 * 100 * 200, "only the last three axes count")
    check(displayed_voxels((64, 64)) == 64 * 64, "2D shapes handled")

    levels = [_Level((60, 1024, 1024)), _Level((60, 512, 512)), _Level((60, 256, 256))]
    check(choose_level(levels, 128_000_000) == 0, "a 63M-voxel volume uses full resolution")
    check(choose_level(levels, 20_000_000) == 1, "a tighter budget steps down one level")
    check(choose_level(levels, 1_000) == 2, "an impossible budget falls back to the coarsest level")
    check(choose_level([_Level((10, 10, 10))], 1) == 0, "a single-level pyramid is always level 0")


def test_roi_names_are_unique() -> None:
    """napari seeds each new shape from the previous one's features."""
    print("ROI names stay unique")
    from microscopy_viewer.measurements import NAME_FEATURE, ensure_name_feature, roi_names

    class _Layer:
        def __init__(self, count, names):
            self.data = [None] * count
            self.features = {NAME_FEATURE: np.array(names, dtype=object)} if names else {}

    # Three shapes that all inherited the default name.
    layer = _Layer(3, ["ROI 1", "ROI 1", "ROI 1"])
    ensure_name_feature(layer)
    names = roi_names(layer)
    check(names == ["ROI 1", "ROI 2", "ROI 3"], f"copied defaults are renumbered: {names}")
    check(len(set(names)) == 3, "all names distinct")

    # Genuine user names survive; only the repeat is replaced.
    layer = _Layer(3, ["Nucleus", "Nucleus", ""])
    ensure_name_feature(layer)
    names = roi_names(layer)
    check(names[0] == "Nucleus", f"the first deliberate name is kept ({names[0]})")
    check(names[1] != "Nucleus" and names[2] != "", f"the duplicate and the blank are filled: {names}")
    check(len(set(names)) == 3, f"all distinct: {names}")

    # Shapes added with no feature column at all.
    layer = _Layer(2, [])
    ensure_name_feature(layer)
    check(roi_names(layer) == ["ROI 1", "ROI 2"], f"defaults created from scratch: {roi_names(layer)}")


def test_excel_export() -> None:
    print("Excel export")
    from microscopy_viewer.metadata import AcquisitionMetadata

    meta = AcquisitionMetadata(image_name="unit", pixel_size_x_um=0.5, pixel_size_y_um=0.5)
    meta.channel(0).name = "GFP"
    rows = [
        Measurement(
            roi_name="ROI 1",
            measurement_type="Area",
            value=50.0,
            unit=MICRON_SQ,
            image_name="unit",
            channel="GFP",
            timestamp="2026-07-27T10:00:00",
            shape_type="rectangle",
            roi_layer="ROIs",
        )
    ]
    frame = measurements_dataframe(rows)
    check("ROI name" in frame.columns, f"columns renamed: {list(frame.columns)[:4]}")
    check(len(frame) == 1, f"one row ({len(frame)})")

    with tempfile.TemporaryDirectory() as directory:
        path = export_measurements(rows, Path(directory) / "out.xlsx", [meta])
        check(path.exists() and path.stat().st_size > 0, f"workbook written ({path.stat().st_size} bytes)")

        import openpyxl

        book = openpyxl.load_workbook(path)
        check(book.sheetnames == ["Measurements", "Acquisition"], f"sheets {book.sheetnames}")
        sheet = book["Measurements"]
        check(sheet.cell(row=2, column=1).value == "ROI 1", "first data row holds the ROI name")
        check(sheet.cell(row=2, column=3).value == 50.0, "value written as a number")


def test_load_errors() -> None:
    print("error handling")
    from microscopy_viewer.loaders import LoadError, load_paths

    specs, errors = load_paths([SAMPLES / "does_not_exist.ims", SAMPLES / "sample_plain.tif"])
    check(len(specs) == 1, f"the good file still loaded ({len(specs)} layer)")
    check(len(errors) == 1 and isinstance(errors[0], LoadError), "the missing file produced one LoadError")

    with tempfile.TemporaryDirectory() as directory:
        bogus = Path(directory) / "bogus.ims"
        bogus.write_bytes(b"not hdf5 at all")
        _specs, errors = load_paths([bogus])
        check(len(errors) == 1, "a corrupt .ims is reported, not raised")


def main() -> int:
    if not SAMPLES.exists():
        print(f"sample data missing — run: python {Path('tests/make_sample_data.py')}")
        return 2
    for test in (
        test_ims_4d,
        test_ims_2d,
        test_brightfield_is_gray,
        test_brightfield_name_matching,
        test_pyramid_padding_without_imagesize_attrs,
        test_ome_tiff,
        test_imagej_tiff,
        test_plain_tiff,
        test_ome_zarr,
        test_ome_zarr_plate,
        test_ome_zarr_acquisitions_are_channels,
        test_ome_zarr_sparse_plate_is_trimmed,
        test_roi_layer_text_matches_the_layer_ndim,
        test_ome_zarr_well_and_descent,
        test_bioformats2raw_series,
        test_measurements,
        test_ellipse_and_uncalibrated,
        test_3d_line_uses_z,
        test_3d_level_choice,
        test_roi_names_are_unique,
        test_excel_export,
        test_load_errors,
    ):
        test()
        print()

    if _failures:
        print(f"{len(_failures)} check(s) failed:")
        for failure in _failures:
            print(f"  - {failure}")
        return 1
    print("all checks passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
