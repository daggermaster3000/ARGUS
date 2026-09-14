"""Checks for overview detection, stitching and the locator slide. No display required.

Run with::

    python tests/test_overview.py
"""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

import numpy as np

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from microscopy_viewer import overview as ov  # noqa: E402
from microscopy_viewer import slides  # noqa: E402

_failures: list[str] = []


def check(condition: bool, message: str) -> None:
    if condition:
        print(f"  ok   {message}")
    else:
        print(f"  FAIL {message}")
        _failures.append(message)


#: A 5x5 grid of 600 µm fields stepping 500 µm, so neighbours overlap by 100 µm —
#: the same one-sixth overlap a real Imaris mosaic is acquired with.
_FIELD = 600.0
_STEP = 500.0


def _tile_sample(index: int, plane: np.ndarray | None = None) -> slides.SampleSlide:
    """One overview field, named and placed the way Imaris writes them."""
    column, row = index % 5, index // 5
    x0, y0 = column * _STEP, row * _STEP
    data = np.zeros((32, 32), dtype=np.uint16) if plane is None else plane
    channel = slides.ChannelView(
        label="Brightfield", color=(1, 1, 1), layer_name=f"tile {index}", data=data,
        axes="YX", channel_index=0,
    )
    return slides.SampleSlide(
        name=f"overview_F{index:04d}",
        channels=[channel],
        pixel_size_um=_FIELD / 32,
        source=f"C:/data/overview_F{index:04d}.ims",
        stage_extent=(x0, x0 + _FIELD, y0, y0 + _FIELD, 0.0, 1.0),
    )


def _acquisition(
    name: str, x: float, y: float, size: float = 120.0, objective: str = ""
) -> slides.SampleSlide:
    """A z-stack sample centred on a stage position."""
    rng = np.random.default_rng(abs(hash(name)) % 1000)
    channels = [
        slides.ChannelView(
            label=f"Ch{index}", color=(0.0, 1.0, 0.0), layer_name=f"{name} Ch{index}",
            data=rng.integers(0, 500, size=(3, 16, 16)).astype(np.uint16),
            axes="ZYX", contrast_limits=(0.0, 500.0), channel_index=index,
        )
        for index in range(2)
    ]
    return slides.SampleSlide(
        name=name, channels=channels, pixel_size_um=0.6, source=f"C:/data/{name}.ims",
        stage_extent=(x - size / 2, x + size / 2, y - size / 2, y + size / 2, 0.0, 50.0),
        objective=objective,
    )


# ---------------------------------------------------------------------------


def test_detection() -> None:
    print("overview fields are told apart from samples")
    tiles = [_tile_sample(index) for index in range(25)]
    samples = [_acquisition("fish_1", 800, 700), _acquisition("fish_2", 1500, 1200)]

    rest, found = ov.split_tiles(tiles + samples)
    check(len(found) == 25 and len(rest) == 2, f"25 fields split from 2 samples ({len(found)}, {len(rest)})")
    check([s.name for s in rest] == ["fish_1", "fish_2"], "the samples keep their order")
    check([ov.field_index(t.name) for t in found] == list(range(25)), "fields come back in field order")

    # A z-stack named like a field is an acquisition, not an overview: Imaris
    # numbers multi-position experiments the same way.
    stack = _acquisition("experiment_F0001", 100, 100)
    rest, found = ov.split_tiles([_tile_sample(0), _tile_sample(1), stack])
    check(len(found) == 2 and rest == [stack], "a z-stack with a field name stays a sample")

    # One field on its own is a file that happens to be named like a field.
    rest, found = ov.split_tiles([_tile_sample(0), *samples])
    check(not found and len(rest) == 3, "a lone field is not an overview")

    # The aborted-acquisition stub Imaris leaves behind is "_F0", one digit.
    check(ov.field_base("curved-SV2_F0") is None, "a single-digit _F0 stub is not a field")
    check(ov.field_base("overview_F0007") == "overview", "the series name is recovered")
    check(ov.field_base("straight-SV2_4") is None, "an ordinary repeat number is not a field")

    # No stage coordinates means nothing can be placed, so nothing is taken away.
    blind = _tile_sample(0)
    blind.stage_extent = None
    rest, found = ov.split_tiles([blind, _tile_sample(1)])
    check(not found, "fields without stage coordinates stay in the figure")


def test_geometry() -> None:
    print("stage coordinates map onto the mosaic")
    box = ov.Box(0.0, 100.0, 200.0, 400.0)
    check(box.centre == (50.0, 300.0), f"the centre is the middle ({box.centre})")
    check(ov.box_from_extent((10, 0, 40, 20, 1, 2)) == ov.Box(0, 10, 20, 40), "reversed extents are ordered")
    check(ov.box_from_extent(None) is None, "no extent, no box")

    outer = ov.union([ov.Box(0, 10, 0, 10), ov.Box(-5, 3, 4, 20)])
    check(outer == ov.Box(-5, 10, 0, 20), f"the union spans both ({outer})")
    check(outer.contains(ov.Box(0, 1, 1, 2)), "containment is inclusive")
    check(not outer.contains(ov.Box(9, 11, 1, 2)), "and catches an overhang")


def test_stitch_places_tiles() -> None:
    print("fields land where the stage says they were")
    # Each field is uniform, with a bright square in its own corner of the grid;
    # if placement were wrong the bright patches would not form a diagonal.
    tiles = []
    for index in range(25):
        plane = np.full((32, 32), 100, dtype=np.uint16)
        if index % 6 == 0:  # the diagonal of a 5x5 grid
            plane[8:24, 8:24] = 4000
        tiles.append(_tile_sample(index, plane))

    mosaic = slides.build_mosaic(tiles, max_pixels=200)
    check(mosaic is not None, "a mosaic is produced")
    check(mosaic.tiles == 25, f"every field contributed ({mosaic.tiles})")
    check(mosaic.image.dtype == np.uint8 and mosaic.image.ndim == 3, "uint8 RGB comes out")

    span = 4 * _STEP + _FIELD  # 2600 µm across, square grid
    check(abs(mosaic.box.width - span) < 1e-6, f"the canvas spans the fields ({mosaic.box.width})")
    check(abs(mosaic.aspect - 1.0) < 0.02, f"a square grid gives a square mosaic ({mosaic.aspect:.3f})")
    check(abs(mosaic.um_per_px - span / 200) < 0.2, f"the scale is reported ({mosaic.um_per_px:.1f} µm/px)")

    grey = mosaic.image[..., 0].astype(np.float32)
    rows, columns = np.nonzero(grey > 200)
    # The bright squares run bottom-left to top-right in stage coordinates, and
    # row 0 is the low-Y edge, so they run top-left to bottom-right in the image.
    correlation = np.corrcoef(rows, columns)[0, 1]
    check(correlation > 0.9, f"the bright fields form the expected diagonal ({correlation:.2f})")

    centre = mosaic.fraction(*ov.Box(0, span, 0, span).centre)
    check(all(abs(value - 0.5) < 0.01 for value in centre), f"the middle maps to the middle ({centre})")


def test_overlap_blends() -> None:
    print("neighbouring fields blend across the overlap")
    # Two uniform fields of different brightness, overlapping by 100 µm. Pasting
    # one over the other leaves a step in the middle of the overlap; a feathered
    # average has to cross from one level to the other without a line on either
    # side of it, which is what a seam looks like.
    dim = ov.Tile(box=ov.Box(0, 600, 0, 600), read=lambda: np.full((64, 64), 100, np.uint16))
    bright = ov.Tile(box=ov.Box(500, 1100, 0, 600), read=lambda: np.full((64, 64), 140, np.uint16))

    mosaic = ov.stitch([dim, bright], max_pixels=220)
    profile = mosaic.image[mosaic.image.shape[0] // 2, :, 0].astype(np.int16)
    check(profile[0] == 0 and profile[-1] == 255, f"each field keeps its own level ({profile[0]}, {profile[-1]})")

    steps = np.diff(profile)
    check(int(steps.min()) >= 0, f"the crossing never reverses ({int(steps.min())})")
    check(int(steps.max()) < 60, f"and never jumps the whole way in one pixel ({int(steps.max())})")
    check(int(np.count_nonzero(steps)) > 5, f"it is a ramp, not a hard edge ({int(np.count_nonzero(steps))} px)")


def test_canvas_covers_outside_samples() -> None:
    print("a sample imaged past the edge of the overview")
    tiles = [_tile_sample(index) for index in range(25)]
    inside = _acquisition("inside", 1000, 1000)
    outside = _acquisition("outside", 4000, 1000)

    mosaic = slides.build_mosaic(tiles, [inside, outside], max_pixels=200)
    check(mosaic.box.x1 >= 4060, f"the canvas is widened to hold it ({mosaic.box.x1:.0f} µm)")
    check(abs(mosaic.covered.x1 - 2600) < 1e-6, "the covered area still records where the fields end")

    marks = ov.markers([inside, outside], mosaic.covered)
    check([m.number for m in marks] == [1, 2], "markers are numbered in deck order")
    check(not marks[0].outside and marks[1].outside, "the one beyond the fields is flagged")

    left, top, width, height = mosaic.fractions_of(marks[0].box)
    check(0 < left < 1 and 0 < top < 1, f"an inside marker is on the picture ({left:.2f}, {top:.2f})")
    check(0 < width < 0.1, f"a 120 µm field is a small box on a 4 mm canvas ({width:.3f})")

    # The uncovered strip is white, not black: it is nothing, not empty data.
    right_edge = mosaic.image[:, -3:, 0]
    check(int(right_edge.min()) == 255, "the part no field reached is left white")

    # The scale bar has to be burned into the fields, not the blank extension, so
    # the view it is drawn on must exclude the widened part.
    covered_view = mosaic.view_of(mosaic.covered)
    check(covered_view.shape[1] < mosaic.image.shape[1], "the covered view is narrower than the canvas")
    check(int(covered_view.max()) < 255, "and it holds image, not blank paper")
    covered_view[:] = 0
    check(int(mosaic.image[:, -3:, 0].min()) == 255, "writing to it leaves the blank margin alone")
    check(int(mosaic.image.min()) == 0, "and does reach the picture itself")

    # Numbering follows the deck even when a sample has no coordinates at all.
    lost = _acquisition("lost", 0, 0)
    lost.stage_extent = None
    numbered = ov.markers([inside, lost, outside], mosaic.covered)
    check([m.number for m in numbered] == [1, 3], f"an unplaceable sample keeps its number free ({[m.number for m in numbered]})")


def test_overview_slide() -> None:
    print("the locator slide")
    try:
        from pptx import Presentation
    except ImportError:
        print("  skip python-pptx is not installed")
        return

    tiles = [_tile_sample(index) for index in range(25)]
    samples = [_acquisition("fish_1", 800, 700), _acquisition("fish_2", 1500, 1200)]

    with tempfile.TemporaryDirectory() as directory:
        # Closeups off: this test is about the locator slide, and with an overview
        # to sit against every sample here would get one of its own as well.
        path = slides.export_slide(
            samples, Path(directory) / "deck.pptx", title="Fish", max_pixels=32,
            overview_tiles=tiles, overview_pixels=300, zoom_slides=False,
        )
        presentation = Presentation(str(path))
        check(len(presentation.slides) == 2, f"the overview leads the deck ({len(presentation.slides)})")

        first = presentation.slides[0]
        pictures = [shape for shape in first.shapes if shape.shape_type == 13]
        check(len(pictures) == 1, f"one stitched picture, not 25 ({len(pictures)})")

        text = " ".join(shape.text_frame.text for shape in first.shapes if shape.has_text_frame)
        check("overview" in text.lower(), "the slide says what it is")
        check("25 overview fields" in text, f"the field count is stated ({text[:60]!r})")
        check("1  fish_1" in text and "2  fish_2" in text, "the legend numbers every sample")
        # The stage coordinates are printed as well as marked: they are what gets
        # typed back into the microscope to find the specimen again.
        check("X 0.80, Y 0.70 mm" in text, f"the legend states where the stage was ({text[-90:]!r})")

        outlines = [shape for shape in first.shapes if shape.shape_type == 1]  # autoshape
        check(len(outlines) == 2, f"one marker box per sample ({len(outlines)})")
        check(
            all(shape.line.color.rgb is not None for shape in outlines),
            "the markers are outlined in the marker colour",
        )

        # The second slide is the ordinary figure, unchanged by any of this.
        table = [shape.table for shape in presentation.slides[1].shapes if shape.has_table][0]
        check(len(table.rows) == 3, f"a header row plus the two samples ({len(table.rows)})")

        # Markers must sit inside the picture, or they are pointing at nothing.
        picture = pictures[0]
        for shape in outlines:
            inside = (
                picture.left <= shape.left and shape.left + shape.width <= picture.left + picture.width
                and picture.top <= shape.top and shape.top + shape.height <= picture.top + picture.height
            )
            check(inside, f"marker at ({shape.left}, {shape.top}) sits on the overview")

    # Without tiles the deck is exactly what it was before this existed.
    with tempfile.TemporaryDirectory() as directory:
        plain = slides.export_slide(samples, Path(directory) / "plain.pptx", max_pixels=32)
        check(len(Presentation(str(plain)).slides) == 1, "no overview, no extra slide")


def test_stitch_failures_are_survivable() -> None:
    print("nothing to stitch")
    check(ov.stitch([]) is None, "no fields gives no mosaic")
    check(slides.build_mosaic([]) is None, "and the caller gets None rather than an exception")

    # A field whose read fails to produce a plane is skipped, not fatal.
    tiles = [_tile_sample(0), _tile_sample(1)]
    broken = ov.Tile(box=ov.Box(0, 600, 0, 600), read=lambda: np.zeros((0, 0)), name="broken")
    good = ov.Tile(box=ov.Box(500, 1100, 0, 600), read=lambda: np.full((32, 32), 500, np.uint16))
    mosaic = ov.stitch([broken, good], max_pixels=100)
    check(mosaic is not None and mosaic.tiles == 2, "one unreadable field does not lose the mosaic")

    # Every field unreadable is a failure, reported as None.
    check(ov.stitch([broken], max_pixels=100) is None, "an entirely unreadable overview gives None")

    # Degenerate coordinates cannot divide by zero.
    flat = ov.Tile(box=ov.Box(0, 0, 0, 0), read=lambda: np.zeros((4, 4), np.uint16))
    check(ov.stitch([flat]) is None, "a zero-sized field is refused")

    _rest, found = ov.split_tiles(tiles)
    check(len(found) == 2, "the fixtures themselves still detect as an overview")


def test_closeup_detection() -> None:
    print("a closeup is paired with the field it was taken from")
    wide = _acquisition("fish_1 10x", 800, 700, size=1200, objective="10x")
    middle = _acquisition("fish_1 20x", 800, 700, size=600, objective="20x")
    tight = _acquisition("fish_1 40x", 700, 620, size=300, objective="40x")
    elsewhere = _acquisition("fish_2 20x", 1500, 1200, size=600, objective="20x")

    order = [wide, middle, tight, elsewhere]
    pairs = {pair.child.name: pair for pair in ov.closeups(order)}
    check(pairs["fish_1 40x"].parent is middle, "the 40x is shown against the 20x, not the 10x")
    check(pairs["fish_1 20x"].parent is wide, "and the 20x against the 10x")
    check("fish_1 10x" not in pairs, "nothing contains the 10x, so it has no context")
    check("fish_2 20x" not in pairs, "a field somewhere else is not a closeup of anything")

    # Two acquisitions of the same field contain each other; neither is a closeup.
    twin = _acquisition("fish_1 20x again", 800, 700, size=600, objective="20x")
    again = {pair.child.name for pair in ov.closeups([middle, twin])}
    check(not again, f"same-sized fields do not pair with each other ({again})")

    # The stage repeats to a few micrometres, so a closeup may sit a hair over the
    # edge of the field it was picked from and still have come from it.
    over = _acquisition("fish_1 40x edge", 800 + 155, 700, size=300, objective="40x")
    edged = {pair.child.name for pair in ov.closeups([middle, over])}
    check("fish_1 40x edge" in edged, f"a few µm of overhang is still a closeup ({edged})")
    far = _acquisition("fish_1 40x off", 800 + 300, 700, size=300, objective="40x")
    check(not ov.closeups([middle, far]), "half outside the field is not")

    # With an overview to fall back on, a sample nothing contains gets a window of
    # it rather than nothing: smaller than the mosaic, and centred on the sample.
    covered = ov.Box(0, 12000, 0, 12000)
    fallback = {pair.child.name: pair for pair in ov.closeups(order, covered)}
    lonely = fallback["fish_1 10x"]
    check(lonely.parent is None, "the overview stands in when no sample contains a field")
    check(lonely.on_overview, "and says so")
    window = lonely.parent_box
    check(window.width < covered.width, f"the window is a part of the overview ({window.width:.0f} µm)")
    check(window.contains(lonely.box), "and it holds the sample it is a window on")
    check(covered.contains(window), f"without running off it ({window.x0:.0f}-{window.x1:.0f} µm)")

    # The box drawn on the context has to land where the stage says it was.
    left, top, width, height = ov.fractions_within(
        pairs["fish_1 40x"].box, pairs["fish_1 40x"].parent_box
    )
    check(
        abs(width - 0.5) < 1e-6 and abs(height - 0.5) < 1e-6,
        f"a 300 µm field is half of a 600 µm one ({width:.3f} x {height:.3f})",
    )
    check(
        abs(left - 50 / 600) < 1e-6 and abs(top - 70 / 600) < 1e-6,
        f"and it sits where it was imaged ({left:.3f}, {top:.3f})",
    )


def test_closeup_slide() -> None:
    print("the closeup slide")
    try:
        from pptx import Presentation
    except ImportError:
        print("  skip python-pptx is not installed")
        return

    wide = _acquisition("fish_1 10x", 800, 700, size=1200, objective="10x")
    tight = _acquisition("fish_1 40x", 700, 620, size=300, objective="40x")

    with tempfile.TemporaryDirectory() as directory:
        path = slides.export_slide(
            [wide, tight], Path(directory) / "deck.pptx", title="Fish", max_pixels=32
        )
        presentation = Presentation(str(path))
        check(
            len(presentation.slides) == 2,
            f"the table, then one closeup slide ({len(presentation.slides)})",
        )

        last = presentation.slides[-1]
        text = " ".join(shape.text_frame.text for shape in last.shapes if shape.has_text_frame)
        check("fish_1 40x" in text and "inside" in text, f"the slide names both ({text[:70]!r})")
        check("10x" in text and "40x" in text, "both magnifications are stated")
        check("X 0.70, Y 0.62 mm" in text, f"so is the stage position ({text[-90:]!r})")

        pictures = [shape for shape in last.shapes if shape.shape_type == 13]
        # The context, plus the closeup's two channels and their merge.
        check(len(pictures) == 4, f"the context and the closeup's panels ({len(pictures)})")

        outlines = [shape for shape in last.shapes if shape.shape_type == 1]
        check(len(outlines) == 1, f"one region box ({len(outlines)})")

        context = min(pictures, key=lambda shape: shape.left)
        box = outlines[0]
        inside = (
            context.left <= box.left and box.left + box.width <= context.left + context.width
            and context.top <= box.top and box.top + box.height <= context.top + context.height
        )
        check(inside, "the box sits on the context picture")
        # The 40x was taken up and left of the 10x centre, and the box has to
        # follow that or it is decoration rather than a locator.
        check(
            box.left + box.width / 2 < context.left + context.width / 2,
            "on the side of it the closeup came from",
        )

    # Turned off, the deck is what it was before any of this existed.
    with tempfile.TemporaryDirectory() as directory:
        plain = slides.export_slide(
            [wide, tight], Path(directory) / "plain.pptx", max_pixels=32, zoom_slides=False
        )
        check(len(Presentation(str(plain)).slides) == 1, "no closeup slides when they are not asked for")


def main() -> int:
    for test in (
        test_detection,
        test_geometry,
        test_stitch_places_tiles,
        test_overlap_blends,
        test_canvas_covers_outside_samples,
        test_overview_slide,
        test_closeup_detection,
        test_closeup_slide,
        test_stitch_failures_are_survivable,
    ):
        test()
        print()

    if _failures:
        print(f"{len(_failures)} check(s) failed:")
        for failure in _failures:
            print(f"  - {failure}")
        return 1
    print("all checks passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
