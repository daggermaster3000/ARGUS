"""Checks for the PowerPoint figure slide. No display required.

Run with::

    python tests/test_slides.py
"""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

import numpy as np

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from microscopy_viewer import slides  # noqa: E402

_failures: list[str] = []


def check(condition: bool, message: str) -> None:
    if condition:
        print(f"  ok   {message}")
    else:
        print(f"  FAIL {message}")
        _failures.append(message)


class _Colormap:
    """A black-to-colour ramp with napari's ``map`` interface."""

    def __init__(self, color):
        self.color = np.asarray(color, dtype=np.float32)

    def map(self, values):
        values = np.asarray(values, dtype=np.float32).reshape(-1, 1)
        rgb = values * self.color[None, :]
        return np.concatenate([rgb, np.ones_like(values)], axis=1)


def _stack(seed: int = 0, shape=(5, 40, 60)) -> np.ndarray:
    rng = np.random.default_rng(seed)
    return rng.integers(0, 1000, size=shape).astype(np.uint16)


def _sample(name="Sample A", n_channels=2, pixel_size=0.2) -> slides.SampleSlide:
    colors = [(0.0, 1.0, 0.0), (1.0, 0.0, 1.0), (0.0, 1.0, 1.0)]
    channels = [
        slides.ChannelView(
            label=f"Ch{index}",
            color=colors[index % len(colors)],
            layer_name=f"{name} - Ch{index}",
            data=_stack(seed=index),
            axes="ZYX",
            contrast_limits=(0.0, 1000.0),
            colormap=_Colormap(colors[index % len(colors)]),
            channel_index=index,
        )
        for index in range(n_channels)
    ]
    return slides.SampleSlide(name=name, channels=channels, pixel_size_um=pixel_size)


# ---------------------------------------------------------------------------


def test_label_colours() -> None:
    print("legend colours stay readable")
    check(slides.text_color((0.0, 1.0, 0.0))[1] > 0, "green stays green")
    white = slides.text_color((1.0, 1.0, 1.0))
    check(max(white) < 200, f"white is darkened so it shows on a white slide ({white})")
    yellow = slides.text_color((1.0, 1.0, 0.0))
    check(max(yellow) < 255 and yellow[2] == 0, f"yellow darkens but keeps its hue ({yellow})")
    check(slides.text_color((0.0, 0.0, 0.0)) == (0, 0, 0), "black is left alone")

    # A dim colour must not be brightened: only over-bright ones are touched.
    dim = slides.text_color((0.2, 0.0, 0.0))
    check(dim == (51, 0, 0), f"a dark colour passes through unchanged ({dim})")

    check(
        slides.colormap_color(_Colormap((0.0, 1.0, 0.0))) == (0.0, 1.0, 0.0),
        "the colormap's top colour is what the label uses",
    )
    check(slides.colormap_color(None) == (1.0, 1.0, 1.0), "no colormap falls back to white")


def test_projection_and_contrast() -> None:
    print("planes are projected and stretched")
    data = np.zeros((4, 10, 10), dtype=np.uint16)
    data[2, 5, 5] = 800  # only visible if Z is actually projected
    channel = slides.ChannelView(
        label="Ch0", color=(0, 1, 0), layer_name="L", data=data, axes="ZYX",
        contrast_limits=(0.0, 800.0), channel_index=0,
    )

    plane, _description = slides.normalized_plane(channel, "Maximum projection")
    check(plane.shape == (10, 10), f"a z-stack flattens to 2D ({plane.shape})")
    check(plane[5, 5] == 1.0, "the brightest voxel survives the projection")

    slice_plane, _ = slides.normalized_plane(channel, "Current slice")
    check(slice_plane.max() == 0.0, "the current slice (z=0) is empty, as it should be")

    mean_plane, _ = slides.normalized_plane(channel, "Mean projection")
    check(abs(float(mean_plane[5, 5]) - 0.25) < 1e-6, "the mean projection averages over Z")

    # Contrast limits, not the data range, set the mapping.
    channel.contrast_limits = (0.0, 1600.0)
    half, _ = slides.normalized_plane(channel, "Maximum projection")
    check(abs(float(half[5, 5]) - 0.5) < 1e-6, "widening the contrast limits dims the pixel")

    flat = slides.ChannelView(
        label="flat", color=(1, 1, 1), layer_name="L", data=np.zeros((4, 4), np.uint16),
        contrast_limits=(5.0, 5.0),
    )
    empty, _ = slides.normalized_plane(flat, "Current slice")
    check(float(empty.max()) == 0.0, "a degenerate contrast range does not divide by zero")


def test_downsample() -> None:
    print("images are shrunk for the deck")
    plane = np.linspace(0, 1, 400 * 800, dtype=np.float32).reshape(400, 800)
    small = slides.downsample(plane, 200)
    check(max(small.shape) == 200, f"the longest edge hits the target ({small.shape})")
    check(abs(small.shape[1] / small.shape[0] - 2.0) < 0.05, "the aspect ratio is kept")
    check(slides.downsample(plane, 4000).shape == plane.shape, "a small image is left alone")

    # Downsampling must not invent brightness: a constant image stays constant.
    constant = np.full((300, 300), 0.4, dtype=np.float32)
    check(
        abs(float(slides.downsample(constant, 50).mean()) - 0.4) < 1e-3,
        "a uniform image keeps its level",
    )


def test_render_sample() -> None:
    print("channel images and the merge")
    sample = _sample(n_channels=2)
    images, merge, description, bar = slides.render_sample(sample, "Maximum projection", max_pixels=32)

    check(len(images) == 2, f"one image per channel ({len(images)})")
    check(all(image.dtype == np.uint8 and image.ndim == 3 for image in images), "uint8 RGB")
    check(images[0].shape[:2] == images[1].shape[:2], "channels come out the same size")
    check(merge is not None and merge.shape == images[0].shape, "the merge matches them")
    check("max over Z" in description, f"the plane is described ({description})")
    check(bar.endswith("µm"), f"a scale bar length is reported ({bar})")

    # Colour purity has to be checked without the scale bar: it is burned in as
    # white pixels, so it puts every component into every image on purpose.
    plain, plain_merge, _description, _bar = slides.render_sample(
        sample, "Maximum projection", max_pixels=32, scale_bar=False
    )
    check(plain[0][..., 0].max() == 0, "the green channel has no red in it")
    check(plain[1][..., 1].max() == 0, "the magenta channel has no green in it")
    # The merge is additive, so it carries both.
    check(
        plain_merge[..., 0].max() > 0 and plain_merge[..., 1].max() > 0,
        "the merge carries both channels",
    )

    # A channel excluded from the merge must not appear in it.
    sample.channels[1].in_merge = False
    _images, merge_only_green, _description, _bar = slides.render_sample(
        sample, "Maximum projection", max_pixels=32, scale_bar=False
    )
    check(merge_only_green[..., 0].max() == 0, "an excluded channel is absent from the merge")

    # No scale bar when the pixel size is unknown.
    sample.pixel_size_um = None
    _i, _m, _d, missing = slides.render_sample(sample, "Maximum projection", max_pixels=32)
    check(missing == "", "uncalibrated data gets no scale bar")


def test_pyramid_level_choice() -> None:
    print("coarse pyramid levels are read when they are enough")
    levels = [
        np.zeros((3, 2040, 2040), np.uint16),
        np.zeros((3, 1020, 1020), np.uint16),
        np.zeros((3, 510, 510), np.uint16),
        np.zeros((3, 255, 255), np.uint16),
    ]
    channel = slides.ChannelView(
        label="Ch0", color=(0, 1, 0), layer_name="L", data=levels[0], axes="ZYX", levels=levels
    )

    check(channel.level_for(900).shape[-1] == 1020, "a 900 px export reads the 1020 level, not 2040")
    check(channel.level_for(1200).shape[-1] == 2040, "asking for more than 1020 goes back to full resolution")
    check(channel.level_for(255).shape[-1] == 255, "a small export reads a much coarser level")
    check(channel.level_for(100).shape[-1] == 255, "it never goes below the coarsest level")
    check(channel.full_width == 2040, "the pixel size still refers to the finest level")

    plain = slides.ChannelView(label="x", color=(1, 1, 1), layer_name="L", data=levels[0])
    check(plain.level_for(900) is levels[0], "a layer with no pyramid uses its only array")

    # A coarser level must not shift the scale bar: it is derived from the finest
    # width, so the bar covers the same physical distance either way.
    sample = slides.SampleSlide(name="S", channels=[channel], pixel_size_um=0.5)
    _images, _merge, _description, coarse_bar = slides.render_sample(sample, "Maximum projection", 900)
    channel.levels = []
    _images, _merge, _description, full_bar = slides.render_sample(sample, "Maximum projection", 900)
    check(coarse_bar == full_bar, f"the scale bar is unchanged by the level read ({coarse_bar})")


def test_sample_naming() -> None:
    print("row labels")

    class _Meta:
        def __init__(self, path, image_name):
            self.file_path = path
            self.image_name = image_name

    # Imaris stores the acquiring machine's own path as the image name.
    meta = _Meta(Path("H:/data/2026-07-27/TEB-BIOTIN_4_deconvolved_F0.ims"), r"D:\Transfer\TEB-BIOTIN_4.ims")
    check(
        slides.sample_name(meta, "layer") == "TEB-BIOTIN_4_deconvolved_F0",
        f"the file stem is used, not the stored path ({slides.sample_name(meta, 'layer')})",
    )

    no_path = _Meta(None, r"D:\Transfer\2026-07-27\NOTEB-CTL.ims")
    check(slides.sample_name(no_path, "layer") == "NOTEB-CTL", "a Windows path in image_name is trimmed too")
    check(slides.sample_name(_Meta(None, ""), "Layer 1") == "Layer 1", "otherwise the layer name is used")
    check(slides.sample_name(None, "Layer 1") == "Layer 1", "no metadata at all still gives a label")


def test_scale_bar() -> None:
    print("scale bar geometry")
    micrometres, pixels = slides.bar_length(0.1, 1000)  # 100 µm across, wants ~20 µm
    check(micrometres == 20 and pixels == 200, f"a round length is picked ({micrometres} µm, {pixels} px)")

    image = np.zeros((100, 200, 3), dtype=np.uint8)
    label = slides.draw_scale_bar(image, 0.5)
    check(label != "" and image.max() == 255, f"a white bar is drawn on a dark image ({label})")
    check(image[:50].max() == 0, "it sits in the lower half")
    check(image[:, :100].max() == 0, "and on the right")

    # A brightfield panel is a pale field: a white bar there is invisible.
    bright = np.full((100, 200, 3), 240, dtype=np.uint8)
    slides.draw_scale_bar(bright, 0.5)
    check(bright.min() == 0, "the bar goes black on a bright image")
    check(int(np.sum(bright == 0)) > 100, "and it is a solid bar, not a stray pixel")

    check(slides.draw_scale_bar(np.zeros((10, 10, 3), np.uint8), 0.0) == "", "no pixel size, no bar")


def test_columns_and_labels() -> None:
    print("channel columns across samples")
    first = _sample("A", n_channels=2)
    second = _sample("B", n_channels=3)
    columns = slides.channel_columns([first, second])
    check(len(columns) == 3, f"the columns are the union across samples ({len(columns)})")
    check([key for key, _label, _color in columns] == ["0", "1", "2"], "keyed by channel index")

    slides.apply_labels([first, second], {"0": "DAPI", "1": "anti-CD31"})
    check(first.channels[0].label == "DAPI", "renaming reaches the first sample")
    check(second.channels[1].label == "anti-CD31", "and every other sample's matching channel")
    check(second.channels[2].label == "Ch2", "channels with no new label keep the original")

    slides.apply_merge_selection([first, second], ["1"])
    check(
        not first.channels[0].in_merge and first.channels[1].in_merge,
        "the merge selection is applied by column",
    )


def test_pptx_output() -> None:
    print("the written presentation")
    try:
        from pptx import Presentation
    except ImportError:
        print("  skip python-pptx is not installed")
        return

    samples = [_sample("Control", n_channels=2), _sample("Treated", n_channels=2)]
    slides.apply_labels(samples, {"0": "DAPI", "1": "anti-CD31"})

    with tempfile.TemporaryDirectory() as directory:
        path = slides.export_slide(
            samples, Path(directory) / "figure", title="Turbidity assay", max_pixels=64
        )
        check(path.suffix == ".pptx" and path.exists(), f"a .pptx is written ({path.name})")

        presentation = Presentation(str(path))
        check(len(presentation.slides) == 1, "one slide")
        slide = presentation.slides[0]

        tables = [shape.table for shape in slide.shapes if shape.has_table]
        check(len(tables) == 1, f"one table ({len(tables)})")
        table = tables[0]
        check(len(table.rows) == 3, f"a header row plus one row per sample ({len(table.rows)})")
        check(len(table.columns) == 4, f"name + two channels + merge ({len(table.columns)})")

        header = [table.cell(0, index).text for index in range(4)]
        check(header == ["Sample", "DAPI", "anti-CD31", "Merge"], f"header reads {header}")
        check(table.cell(1, 0).text.startswith("Control"), "the row is labelled with the sample")
        check("scale bar" in table.cell(1, 0).text, "and notes the scale bar length")

        # The heading must be printed in the channel's own colour.
        run = table.cell(0, 1).text_frame.paragraphs[0].runs[0]
        check(run.font.color.rgb is not None, "the channel heading is coloured")
        # Pure green has luminance 0.7152; held to a 0.55 ceiling that is 0xC4.
        check(str(run.font.color.rgb) == "00C400", f"green, darkened to read on white ({run.font.color.rgb})")

        pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
        check(len(pictures) == 6, f"two channels plus a merge, twice over ({len(pictures)})")

        titles = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
        check(any("Turbidity assay" in text for text in titles), "the title is on the slide")

    with tempfile.TemporaryDirectory() as directory:
        try:
            slides.export_slide([], Path(directory) / "empty.pptx")
            check(False, "an empty export is refused")
        except ValueError:
            check(True, "an empty export is refused")


def test_auto_contrast_mode() -> None:
    print("contrast modes")
    # A dim image with one hot pixel: the whole point of percentile contrast.
    rng = np.random.default_rng(3)
    data = rng.integers(80, 260, size=(1, 40, 40)).astype(np.uint16)
    data[0, 0, 0] = 60000
    channel = slides.ChannelView(
        label="Ch0", color=(0, 1, 0), layer_name="L", data=data, axes="ZYX",
        contrast_limits=(0.0, 60000.0), channel_index=0,
    )

    displayed, _ = slides.normalized_plane(channel, "Maximum projection", None, slides.CONTRAST_AS_DISPLAYED)
    check(float(displayed.mean()) < 0.02, f"as displayed keeps the layer's wide limits ({displayed.mean():.3f})")

    auto, _ = slides.normalized_plane(channel, "Maximum projection", None, slides.CONTRAST_AUTO)
    check(float(auto.mean()) > 0.3, f"auto stretches the bulk of the image ({auto.mean():.2f})")
    check(float(auto.max()) == 1.0, "and clips the hot pixel rather than scaling everything to it")

    low, high = slides.auto_limits(np.asarray([0, 1, 2, 3, 4, 1000], dtype=np.float32))
    check(high < 1000, f"the 99.5th percentile excludes the outlier ({high:.0f})")

    # When the percentiles collapse there is nothing to stretch; falling back to
    # the full range is what the Auto Contrast button does, so it matches.
    flat = slides.auto_limits(np.zeros((4, 4), dtype=np.float32))
    check(flat == (0.0, 0.0), "a constant image gives a degenerate range, handled downstream")


def test_multiple_slides() -> None:
    print("a folder is split across slides")
    try:
        from pptx import Presentation
    except ImportError:
        print("  skip python-pptx is not installed")
        return

    samples = [_sample(f"Fish {index}", n_channels=2) for index in range(9)]
    with tempfile.TemporaryDirectory() as directory:
        path = slides.export_slide(
            samples, Path(directory) / "batch.pptx", max_pixels=32, rows_per_slide=4
        )
        presentation = Presentation(str(path))
        check(len(presentation.slides) == 3, f"nine samples at four a slide is three slides ({len(presentation.slides)})")

        rows = []
        heights = []
        for slide in presentation.slides:
            table = [shape.table for shape in slide.shapes if shape.has_table][0]
            rows.append(len(table.rows) - 1)
            heights.append(table.rows[1].height)
        check(rows == [4, 4, 1], f"the last slide holds the remainder ({rows})")
        check(len(set(heights)) == 1, "every slide uses the same row height, including the short one")

        titles = [
            shape.text_frame.text
            for slide in presentation.slides
            for shape in slide.shapes
            if shape.has_text_frame and not shape.has_table
        ]
        check(any("(1 of 3)" in text for text in titles), f"slides are numbered ({titles[:1]})")

        # Column headings must be identical on every slide.
        headers = []
        for slide in presentation.slides:
            table = [shape.table for shape in slide.shapes if shape.has_table][0]
            headers.append(tuple(table.cell(0, c).text for c in range(len(table.columns))))
        check(len(set(headers)) == 1, f"the columns match across slides ({headers[0]})")

        single = slides.export_slide(
            samples[:3], Path(directory) / "one.pptx", max_pixels=32, rows_per_slide=4
        )
        one = Presentation(str(single))
        check(len(one.slides) == 1, "three samples still fit on one slide")
        text = [
            shape.text_frame.text
            for shape in one.slides[0].shapes
            if shape.has_text_frame and not shape.has_table
        ]
        check(not any("of 1" in item for item in text), f"a single slide is not numbered ({text})")

        # An absurd row count is clamped rather than producing unusable slides.
        clamped = slides.export_slide(
            samples, Path(directory) / "clamped.pptx", max_pixels=32, rows_per_slide=500
        )
        check(
            len(Presentation(str(clamped)).slides) == 1,
            f"rows per slide is capped at {slides.MAX_ROWS_PER_SLIDE}",
        )


def test_cancellation() -> None:
    print("a long batch can be stopped")
    samples = [_sample(f"Fish {index}", n_channels=1) for index in range(6)]
    seen: list[str] = []

    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "cancelled.pptx"
        try:
            slides.export_slide(
                samples, path, max_pixels=32, rows_per_slide=2,
                progress=lambda index, total, name: seen.append(name),
                should_cancel=lambda: len(seen) >= 3,
            )
            check(False, "cancelling raises ExportCancelled")
        except slides.ExportCancelled as exc:
            check(True, f"cancelling raises ExportCancelled ({exc})")
        check(len(seen) == 3, f"it stops promptly, not at the end ({len(seen)} of 6 rendered)")
        check(not path.exists(), "and no half-written deck is left behind")


def test_batch_reads_files() -> None:
    print("batch mode reads files without a viewer")
    samples_dir = Path(__file__).resolve().parent / "sample_data"
    if not samples_dir.exists():
        print("  skip run tests/make_sample_data.py first")
        return

    found, errors = slides.samples_from_paths([samples_dir])
    check(len(found) >= 3, f"the folder yields {len(found)} sample(s)")
    check(not errors, f"nothing failed to read ({errors})")
    check(all(sample.channels for sample in found), "every sample has at least one channel")
    check(
        all(channel.colormap is not None for sample in found for channel in sample.channels),
        "each channel resolved a colormap, so it can be drawn",
    )
    check(
        all(len(sample.name) > 0 and "\\" not in sample.name for sample in found),
        "row labels are plain names, not paths",
    )

    multichannel = [sample for sample in found if len(sample.channels) > 1]
    check(bool(multichannel), f"multichannel files stay grouped as one row ({len(multichannel)})")

    with tempfile.TemporaryDirectory() as directory:
        path = slides.export_slide(
            found, Path(directory) / "folder.pptx", max_pixels=64, rows_per_slide=3,
            contrast=slides.CONTRAST_AUTO,
        )
        check(path.exists(), "the folder exports without the viewer ever being built")

    missing, errors = slides.samples_from_paths([samples_dir / "does_not_exist.ims"])
    check(not missing and len(errors) == 1, f"a bad path is reported, not raised ({errors})")


def test_missing_channel_is_survivable() -> None:
    print("samples with different channel sets")
    try:
        from pptx import Presentation
    except ImportError:
        print("  skip python-pptx is not installed")
        return

    full = _sample("Stained", n_channels=3)
    partial = _sample("Unstained control", n_channels=1)
    with tempfile.TemporaryDirectory() as directory:
        path = slides.export_slide([full, partial], Path(directory) / "gaps.pptx", max_pixels=32)
        table = [shape.table for shape in Presentation(str(path)).slides[0].shapes if shape.has_table][0]
        check(len(table.columns) == 5, f"columns follow the richer sample ({len(table.columns)})")
        check(table.cell(2, 2).text == "—", "a missing channel leaves a marked gap, not a crash")

    # A brightfield-only dataset has nothing in the merge; the cell must say so
    # rather than looking like a panel that failed to render.
    lone = _sample("Brightfield only", n_channels=1)
    lone.channels[0].in_merge = False
    with tempfile.TemporaryDirectory() as directory:
        path = slides.export_slide([lone], Path(directory) / "nomerge.pptx", max_pixels=32)
        table = [shape.table for shape in Presentation(str(path)).slides[0].shapes if shape.has_table][0]
        check(table.cell(1, 2).text == "—", "an empty merge cell is marked too")
        pictures = [s for s in Presentation(str(path)).slides[0].shapes if s.shape_type == 13]
        check(len(pictures) == 1, f"and only the one real panel is placed ({len(pictures)})")


def main() -> int:
    for test in (
        test_label_colours,
        test_projection_and_contrast,
        test_downsample,
        test_render_sample,
        test_pyramid_level_choice,
        test_sample_naming,
        test_scale_bar,
        test_columns_and_labels,
        test_auto_contrast_mode,
        test_pptx_output,
        test_multiple_slides,
        test_cancellation,
        test_batch_reads_files,
        test_missing_channel_is_survivable,
    ):
        test()
        print()

    if _failures:
        print(f"{len(_failures)} check(s) failed:")
        for failure in _failures:
            print(f"  - {failure}")
        return 1
    print("all checks passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
