"""Export the "One image, start to end" diagram of pipeline_explainer.html as SVG, PNG and PPTX.

The diagram is drawn by the page's script, so the page is opened in headless
Chrome and the finished SVG is read off it. From that:

- ``workflow.svg``  — the vector diagram (PowerPoint: Insert ▸ Pictures, then *Convert to Shape*);
- ``workflow.png``  — at twice the size, for slides and documents;
- ``workflow.pptx`` — one slide of native shapes: boxes, pictures, text boxes and arrows,
  each editable on its own.

    python notebooks/explainers/export_workflow.py
"""
from __future__ import annotations

import base64
import io
import re
import subprocess
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path

HERE = Path(__file__).resolve().parent
PAGE = HERE / "pipeline_explainer.html"
OUT = HERE / "workflow"
CHROME = Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe")
SVGNS = "http://www.w3.org/2000/svg"
FONT, MONO = "Segoe UI", "Consolas"
PAPER = "#faf7f2"


def diagram_svg() -> str:
    """The diagram, as the page draws it, made into a standalone SVG file."""
    dom = subprocess.run([str(CHROME), "--headless=new", "--disable-gpu", "--no-sandbox",
                          "--virtual-time-budget=60000", "--window-size=1400,1000", "--dump-dom",
                          PAGE.as_uri() + "?shot=workflow"], capture_output=True, text=True,
                         encoding="utf-8", timeout=600).stdout
    start = dom.index("<svg", dom.index('id="wf"'))
    svg = dom[start: dom.index("</svg>", start) + 6]
    width, height = (float(v) for v in re.search(r'viewBox="0 0 ([\d.]+) ([\d.]+)"', svg).groups())
    svg = svg.replace("<svg ", f'<svg xmlns="{SVGNS}" xmlns:xlink="http://www.w3.org/1999/xlink" '
                               f'width="{width:g}" height="{height:g}" font-family="{FONT}, sans-serif" ', 1)
    svg = svg.replace("var(--mono)", f"{MONO}, monospace")
    svg = re.sub(r"<title>.*?</title>", "", svg, flags=re.S)          # hover texts: page only
    # A paper background, as on the page.
    svg = re.sub(r"(<svg[^>]*>)", rf'\1<rect width="100%" height="100%" fill="{PAPER}"/>', svg, count=1)
    return svg


def to_png(svg_path: Path, png_path: Path, width: float, height: float, scale: int = 2) -> None:
    html = Path(tempfile.mkdtemp()) / "wf.html"
    html.write_text(f'<html><body style="margin:0;background:{PAPER}"><img src="{svg_path.as_uri()}" '
                    f'width="{width:g}" height="{height:g}"></body></html>', encoding="utf-8")
    subprocess.run([str(CHROME), "--headless=new", "--disable-gpu", "--no-sandbox", "--hide-scrollbars",
                    f"--force-device-scale-factor={scale}", f"--window-size={int(width)},{int(height)}",
                    f"--screenshot={png_path}", html.as_uri()], capture_output=True, timeout=300)


# ---------------------------------------------------------------------------
# PPTX from the SVG elements
# ---------------------------------------------------------------------------


def _rgb(colour: str):
    from pptx.dml.color import RGBColor

    colour = colour.strip()
    if colour.startswith("#") and len(colour) == 7:
        return RGBColor.from_string(colour[1:].upper())
    m = re.match(r"rgba?\(([\d.]+),\s*([\d.]+),\s*([\d.]+)", colour)
    if m:
        return RGBColor(*(int(float(v)) for v in m.groups()))
    return RGBColor(0x1F, 0x24, 0x30)


def _points(d: str):
    nums = [float(v) for v in re.findall(r"-?[\d.]+", d)]
    return list(zip(nums[0::2], nums[1::2]))


def _arrow_end(shape) -> None:
    from pptx.oxml.ns import qn

    ln = shape.line._get_or_add_ln()
    tail = ln.find(qn("a:tailEnd"))
    if tail is None:
        tail = ln.makeelement(qn("a:tailEnd"), {})
        ln.append(tail)
    tail.set("type", "triangle"); tail.set("w", "med"); tail.set("len", "med")


def to_pptx(svg: str, path: Path) -> None:
    from PIL import Image
    from pptx import Presentation
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Pt

    root = ET.fromstring(svg)
    width, height = (float(v) for v in root.get("viewBox").split()[2:])
    prs = Presentation()
    prs.slide_width = Emu(int(13.333 * 914400))
    unit = prs.slide_width / width                               # EMU per diagram unit
    prs.slide_height = Emu(int(height * unit))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    E = lambda v: Emu(int(round(float(v) * unit)))               # noqa: E731
    pt = lambda px: Pt(float(px) * unit / 12700)                 # noqa: E731  (1 pt = 12700 EMU)
    tag = lambda el: el.tag.split("}")[-1]                       # noqa: E731

    def rect(el):
        if str(el.get("width", "")).endswith("%"):               # the page background
            slide.background.fill.solid(); slide.background.fill.fore_color.rgb = _rgb(el.get("fill", PAPER))
            return
        x, y, w, h = (float(el.get(k, 0)) for k in ("x", "y", "width", "height"))
        if w >= width - 1 and h >= height - 1:
            slide.background.fill.solid(); slide.background.fill.fore_color.rgb = _rgb(el.get("fill", PAPER))
            return
        rounded = float(el.get("rx", 0)) > 0
        s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, E(x), E(y), E(w), E(h))
        if rounded:
            s.adjustments[0] = min(0.5, float(el.get("rx")) / max(min(w, h), 1))
        fill = el.get("fill", "none")
        if fill == "none":
            s.fill.background()
        else:
            s.fill.solid(); s.fill.fore_color.rgb = _rgb(fill)
        stroke = el.get("stroke")
        if stroke and stroke != "none":
            s.line.color.rgb = _rgb(stroke); s.line.width = pt(el.get("stroke-width", 1))
        else:
            s.line.fill.background()
        s.shadow.inherit = False

    def image(el):
        href = el.get("href") or el.get("{http://www.w3.org/1999/xlink}href")
        data = base64.b64decode(href.split(",", 1)[1])
        x, y, w, h = (float(el.get(k)) for k in ("x", "y", "width", "height"))
        iw, ih = Image.open(io.BytesIO(data)).size
        k = min(w / iw, h / ih)                                  # preserveAspectRatio: meet, centred
        slide.shapes.add_picture(io.BytesIO(data), E(x + (w - iw * k) / 2), E(y + (h - ih * k) / 2), E(iw * k), E(ih * k))

    def text(el):
        size = float(el.get("font-size", 12))
        x, y = float(el.get("x", 0)), float(el.get("y", 0))
        spans = [s for s in el if tag(s) == "tspan"]
        lines = [s.text or "" for s in spans] if spans else [el.text or ""]
        if not any(line.strip() for line in lines):
            return
        anchor = el.get("text-anchor", "start")
        longest = max(len(line) for line in lines)
        w = max(20.0, longest * size * 0.58 + 6)
        h = size * 1.25 * len(lines) + 2
        left = x if anchor == "start" else x - w / 2 if anchor == "middle" else x - w
        top = y - size * 0.95
        rotate = re.search(r"rotate\(([-\d.]+)", el.get("transform", ""))
        if rotate:                                               # turn about the text's own anchor point
            left, top = x - w / 2, y - h / 2
        box = slide.shapes.add_textbox(E(left), E(top), E(w), E(h))
        tf = box.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.TOP
        for k, line in enumerate(lines):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = {"start": PP_ALIGN.LEFT, "middle": PP_ALIGN.CENTER, "end": PP_ALIGN.RIGHT}[anchor]
            run = p.add_run(); run.text = line
            f = run.font
            f.size = pt(size); f.color.rgb = _rgb(el.get("fill", "#1f2430"))
            f.bold = el.get("font-weight", "400") in ("600", "650", "700", "bold")
            f.italic = el.get("font-style") == "italic"
            f.name = MONO if "monospace" in (el.get("font-family") or "") else FONT
        if rotate:
            box.rotation = float(rotate.group(1)) % 360

    def arrow(el):
        if el.get("fill") not in (None, "none"):
            return                                               # the arrowheads inside <marker>
        pts = _points(el.get("d"))
        if len(pts) < 2:
            return
        if len(pts) == 2:
            (x0, y0), (x1, y1) = pts
            s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x0), E(y0), E(x1), E(y1))
        else:
            builder = slide.shapes.build_freeform(E(pts[0][0]), E(pts[0][1]), scale=1.0)
            builder.add_line_segments([(E(a), E(b)) for a, b in pts[1:]], close=False)
            s = builder.convert_to_shape()
            s.fill.background()
        s.line.color.rgb = _rgb(el.get("stroke", "#1f2430"))
        s.line.width = pt(el.get("stroke-width", 2))
        if el.get("stroke-dasharray"):
            s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if el.get("marker-end"):
            _arrow_end(s)

    def walk(el, in_marker=False):
        name = tag(el)
        if name in ("defs", "marker", "title"):
            return
        if name == "rect":
            rect(el)
        elif name == "image":
            image(el)
        elif name == "text":
            text(el)
            return
        elif name == "path":
            arrow(el)
        for child in el:
            walk(child)

    walk(root)
    prs.save(path)


def main() -> None:
    OUT.mkdir(exist_ok=True)
    svg = diagram_svg()
    width, height = (float(v) for v in re.search(r'viewBox="0 0 ([\d.]+) ([\d.]+)"', svg).groups())
    (OUT / "workflow.svg").write_text(svg, encoding="utf-8")
    to_png(OUT / "workflow.svg", OUT / "workflow.png", width, height)
    to_pptx(svg, OUT / "workflow.pptx")
    for name in ("workflow.svg", "workflow.png", "workflow.pptx"):
        print(f"wrote {OUT / name} ({(OUT / name).stat().st_size / 1e6:.2f} MB)")


if __name__ == "__main__":
    main()
